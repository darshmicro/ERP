"""Controlled Excel import (spec 87): Upload -> Validate -> Preview -> Error report -> Submit ->
Approve (e-signature, different user) -> Import. Uploaded data is staged in import_row and only
created through the normal domain services (as DRAFT records) after approval.
"""
import io
import json
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Callable

from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.audit.context import set_reason
from app.core import db
from app.core.errors import BusinessRuleError, NotFound, PermissionDenied, ValidationFailed
from app.core.time import utcnow
from app.models.importing import ImportJob, ImportRow
from app.models.master import Category, Customer, Location, Material, MaterialType, Unit, Vendor, Warehouse
from app.models.spec import STP, Specification, SpecificationParameter
from app.security.permissions import effective_permissions
from app.services import esign, master_services as ms, masters, sod, versioning
from app.services import documents
from app.workflows.state_machine import StateMachine, Transition, transition

MAX_ROWS = 5000
IMPORT_MACHINE = StateMachine("import_job", "UPLOADED", {
    "UPLOADED": {"VALIDATED": Transition("VALIDATED"), "FAILED": Transition("FAILED")},
    "VALIDATED": {"SUBMITTED": Transition("SUBMITTED")},
    "SUBMITTED": {"APPROVED": Transition("APPROVED", "import.job.approve", "APPROVED_BY", True),
                  "REJECTED": Transition("REJECTED", "import.job.approve", "REJECTED_BY", True)},
    "APPROVED": {"IMPORTED": Transition("IMPORTED"), "FAILED": Transition("FAILED")},
})
ENTITY_APPROVE_PERM = {"vendor": "md.vendor.approve", "material": "md.material.approve",
                       "specification": "md.spec.approve", "stp": "md.stp.approve"}

# entity -> (required columns, optional columns)
COLUMNS: dict[str, tuple[list[str], list[str]]] = {
    "vendor": (["name"], ["vendor_code", "vendor_type", "address", "country", "state", "city", "gst_no", "pan_no",
                          "contact_person", "email", "phone", "risk_class", "material_categories"]),
    "material": (["name", "type_code", "base_unit_code"],
                 ["material_code", "category_code", "generic_name", "grade", "pharmacopoeial_standard", "manufacturer",
                  "pack_size", "storage_condition", "temp_min", "temp_max", "shelf_life_days", "retest_days",
                  "gmp_criticality", "hazard_class", "fefo_mode"]),
    "location": (["warehouse_code", "location_code", "name", "location_type"],
                 ["parent_code", "storage_condition", "temp_min", "temp_max", "is_quarantine", "capacity"]),
    "stp": (["title", "procedure"], ["stp_no", "test_method", "equipment_required", "reagents_required",
                                      "reference_standards", "calculation", "acceptance_criteria", "safety_precautions"]),
    "specification": (["material_code", "test_name"],
                      ["spec_type", "lsl", "usl", "target", "unit", "decimal_places", "test_method", "stp_no",
                       "acceptance_criteria", "criticality", "frequency", "pharmacopoeial_reference"]),
}
DESCRIPTIONS = {
    "risk_class": "CRITICAL | HIGH | MEDIUM | LOW", "location_type": "ZONE | ROOM | RACK | SHELF | BIN",
    "fefo_mode": "FEFO | FIFO", "spec_type": "NUMERIC | RANGE | TEXT | PASS_FAIL",
    "criticality": "CRITICAL | MAJOR | MINOR", "gmp_criticality": "CRITICAL | MAJOR | MINOR",
    "is_quarantine": "Yes | No", "type_code": "existing material type code (e.g. RM)",
    "base_unit_code": "existing unit code (e.g. kg)", "material_code": "existing, ACTIVE or APPROVED material code",
    "parent_code": "location code of the parent - must appear earlier in the file or already exist",
}


def template_xlsx(entity: str) -> bytes:
    req, opt = _spec(entity)
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(req + opt)
    ins = wb.create_sheet("Instructions")
    ins.append(["Column", "Required", "Notes"])
    for c in req + opt:
        ins.append([c, "Yes" if c in req else "No", DESCRIPTIONS.get(c, "")])
    ins.append([])
    ins.append(["Rules", "", "Row 1 must be the header. Max 5000 rows. Records are created as DRAFT after QA approval of the import."])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _spec(entity: str) -> tuple[list[str], list[str]]:
    if entity not in COLUMNS:
        raise ValidationFailed(f"Import is not available for '{entity}'")
    return COLUMNS[entity]


def parse_workbook(data: bytes) -> tuple[list[str], list[dict]]:
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValidationFailed("The file is not a valid .xlsx workbook") from exc
    ws = wb["Data"] if "Data" in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    try:
        header = [str(h).strip().lower() if h is not None else "" for h in next(rows)]
    except StopIteration:
        raise ValidationFailed("The workbook is empty")
    out = []
    for r in rows:
        if r is None or all(c is None or str(c).strip() == "" for c in r):
            continue
        out.append({header[i]: _norm(v) for i, v in enumerate(r) if i < len(header) and header[i]})
        if len(out) > MAX_ROWS:
            raise ValidationFailed(f"More than {MAX_ROWS} rows; split the file")
    return header, out


def _norm(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    s = str(v).strip()
    return s or None


# ------------------------------------------------------------ per-row validation
def _dec(v: Any, label: str, errs: list[str]) -> Decimal | None:
    if v is None:
        return None
    try:
        return Decimal(str(v))
    except InvalidOperation:
        errs.append(f"{label} must be a number")
        return None


def _int(v: Any, label: str, errs: list[str]) -> int | None:
    if v is None:
        return None
    try:
        return int(Decimal(str(v)))
    except (InvalidOperation, ValueError):
        errs.append(f"{label} must be a whole number")
        return None


def _bool(v: Any) -> bool:
    return str(v).strip().lower() in ("1", "true", "yes", "y")


def _enum(v: Any, allowed: set[str], label: str, errs: list[str], default: str | None = None) -> str | None:
    if v is None:
        return default
    if str(v).upper() not in allowed:
        errs.append(f"{label} must be one of {sorted(allowed)}")
        return None
    return str(v).upper()


def validate_rows(session: Session, entity: str, rows: list[dict]) -> list[tuple[dict, list[str]]]:
    req, opt = _spec(entity)
    results: list[tuple[dict, list[str]]] = []
    seen: dict[str, int] = {}
    file_locations: set[str] = set()
    for i, row in enumerate(rows, 1):
        errs: list[str] = []
        unknown = set(row) - set(req) - set(opt)
        if unknown:
            errs.append(f"Unknown column(s): {', '.join(sorted(unknown))}")
        for c in req:
            if row.get(c) is None:
                errs.append(f"{c} is required")
        clean = {k: v for k, v in row.items() if k in req + opt}
        if not errs:
            fn = _ROW_VALIDATORS[entity]
            fn(session, clean, errs, seen, i, file_locations)
        results.append((clean, errs))
    return results


def _v_vendor(session, r, errs, seen, i, _):
    key = r["name"].lower()
    if key in seen:
        errs.append(f"Duplicate vendor name (also row {seen[key]})")
    seen[key] = i
    if session.execute(select(Vendor.id).where(func.lower(Vendor.name) == key)).first():
        errs.append("Vendor name already exists")
    if r.get("vendor_code") and session.execute(select(Vendor.id).where(Vendor.vendor_code == r["vendor_code"])).first():
        errs.append("vendor_code already exists")
    r["risk_class"] = _enum(r.get("risk_class"), {"CRITICAL", "HIGH", "MEDIUM", "LOW"}, "risk_class", errs, "MEDIUM")
    if r.get("email") and not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", r["email"]):
        errs.append("email is not valid")
    if r.get("pan_no") and not re.match(r"^[A-Z]{5}[0-9]{4}[A-Z]$", r["pan_no"].upper()):
        errs.append("pan_no format is invalid")


def _v_material(session, r, errs, seen, i, _):
    key = r["name"].lower()
    if key in seen:
        errs.append(f"Duplicate material name (also row {seen[key]})")
    seen[key] = i
    t = session.execute(select(MaterialType).where(func.lower(MaterialType.code) == r["type_code"].lower())).scalar_one_or_none()
    u = session.execute(select(Unit).where(func.lower(Unit.code) == r["base_unit_code"].lower())).scalar_one_or_none()
    if t is None:
        errs.append(f"Unknown type_code '{r['type_code']}'")
    if u is None:
        errs.append(f"Unknown base_unit_code '{r['base_unit_code']}'")
    cat = None
    if r.get("category_code"):
        cat = session.execute(select(Category).where(func.lower(Category.code) == r["category_code"].lower())).scalar_one_or_none()
        if cat is None:
            errs.append(f"Unknown category_code '{r['category_code']}'")
    if r.get("material_code") and session.execute(select(Material.id).where(Material.material_code == r["material_code"])).first():
        errs.append("material_code already exists")
    for f in ("temp_min", "temp_max"):
        r[f] = _dec(r.get(f), f, errs)
    for f in ("shelf_life_days", "retest_days"):
        r[f] = _int(r.get(f), f, errs)
    r["gmp_criticality"] = _enum(r.get("gmp_criticality"), {"CRITICAL", "MAJOR", "MINOR"}, "gmp_criticality", errs, "MAJOR")
    r["fefo_mode"] = _enum(r.get("fefo_mode"), {"FEFO", "FIFO"}, "fefo_mode", errs, "FEFO")
    if not errs:
        r["_type_id"], r["_unit_id"], r["_category_id"] = t.id, u.id, cat.id if cat else None
        try:
            ms._check_ranges({k: r.get(k) for k in ("temp_min", "temp_max", "shelf_life_days", "retest_days")})
        except ValidationFailed as e:
            errs.append(e.message)


def _v_location(session, r, errs, seen, i, file_locations):
    code = r["location_code"]
    if code.lower() in seen:
        errs.append(f"Duplicate location_code (also row {seen[code.lower()]})")
    seen[code.lower()] = i
    if session.execute(select(Location.id).where(Location.location_code == code)).first():
        errs.append("location_code already exists")
    wh = session.execute(select(Warehouse).where(Warehouse.warehouse_code == r["warehouse_code"])).scalar_one_or_none()
    if wh is None:
        errs.append(f"Unknown warehouse_code '{r['warehouse_code']}'")
    lt = _enum(r["location_type"], {"ZONE", "ROOM", "RACK", "SHELF", "BIN"}, "location_type", errs)
    r["location_type"] = lt
    parent_type = None
    if r.get("parent_code"):
        parent = session.execute(select(Location).where(Location.location_code == r["parent_code"])).scalar_one_or_none()
        if parent is not None:
            parent_type = parent.location_type
            if wh is not None and parent.warehouse_id != wh.id:
                errs.append("parent_code belongs to a different warehouse")
        elif r["parent_code"] not in file_locations:
            errs.append(f"parent_code '{r['parent_code']}' not found (must appear earlier in the file)")
        else:
            parent_type = seen.get(("type", r["parent_code"]))
    if lt and parent_type and ms._level(lt) <= ms._level(parent_type):
        errs.append(f"{lt} cannot be under {parent_type}")
    seen[("type", code)] = lt
    file_locations.add(code)
    for f in ("temp_min", "temp_max", "capacity"):
        r[f] = _dec(r.get(f), f, errs)
    r["is_quarantine"] = _bool(r.get("is_quarantine"))


def _v_stp(session, r, errs, seen, i, _):
    if r.get("stp_no"):
        if r["stp_no"].lower() in seen:
            errs.append(f"Duplicate stp_no (also row {seen[r['stp_no'].lower()]})")
        seen[r["stp_no"].lower()] = i
        if session.execute(select(STP.id).where(STP.stp_no == r["stp_no"])).first():
            errs.append("stp_no already exists - create a new version in the STP screen instead")


def _v_spec(session, r, errs, seen, i, _):
    m = session.execute(select(Material).where(Material.material_code == r["material_code"])).scalar_one_or_none()
    if m is None:
        errs.append(f"Unknown material_code '{r['material_code']}'")
    elif m.master_status not in ("APPROVED", "ACTIVE"):
        errs.append(f"Material {m.material_code} is {m.master_status}; it must be APPROVED or ACTIVE")
    r["spec_type"] = _enum(r.get("spec_type"), {"NUMERIC", "RANGE", "TEXT", "PASS_FAIL"}, "spec_type", errs, "NUMERIC")
    r["criticality"] = _enum(r.get("criticality"), {"CRITICAL", "MAJOR", "MINOR"}, "criticality", errs, "MAJOR")
    for f in ("lsl", "usl", "target"):
        r[f] = _dec(r.get(f), f, errs)
    r["decimal_places"] = _int(r.get("decimal_places"), "decimal_places", errs)
    if r["lsl"] is not None and r["usl"] is not None and r["usl"] < r["lsl"]:
        errs.append("usl is below lsl")
    if r.get("spec_type") in ("NUMERIC", "RANGE") and r["lsl"] is None and r["usl"] is None:
        errs.append("numeric tests need lsl and/or usl")
    k = (r["material_code"], r["test_name"].lower())
    if k in seen:
        errs.append(f"Duplicate test for this material (also row {seen[k]})")
    seen[k] = i
    if r.get("stp_no"):
        stp = session.execute(select(STP).where(STP.stp_no == r["stp_no"], STP.status == "APPROVED")).scalars().first()
        if stp is None:
            errs.append(f"stp_no '{r['stp_no']}' has no APPROVED version")
        else:
            r["_stp_id"] = stp.id


_ROW_VALIDATORS: dict[str, Callable] = {"vendor": _v_vendor, "material": _v_material, "location": _v_location,
                                         "stp": _v_stp, "specification": _v_spec}


# ------------------------------------------------------------ job workflow
def upload(session: Session, entity: str, filename: str, data: bytes, user) -> ImportJob:
    _spec(entity)
    if not filename.lower().endswith(".xlsx"):
        raise ValidationFailed("Only .xlsx files can be imported")
    doc = documents.store(session, filename, data, allowed_ext={".xlsx"})
    header, rows = parse_workbook(data)
    req, opt = _spec(entity)
    missing = [c for c in req if c not in header]
    if missing:
        raise ValidationFailed(f"Missing required column(s): {', '.join(missing)}")
    job = ImportJob(entity=entity, filename=filename[:255], document_id=doc.id, created_by_id=user.id,
                    total_rows=len(rows))
    session.add(job)
    session.flush()
    results = validate_rows(session, entity, rows)
    for n, (clean, errs) in enumerate(results, 2):  # row numbers as seen in Excel (header = 1)
        session.add(ImportRow(job_id=job.id, row_no=n, data_json=json.dumps(clean, default=str),
                              status="ERROR" if errs else "VALID", errors_json=json.dumps(errs) if errs else None))
    job.error_rows = sum(1 for _c, e in results if e)
    transition(session, IMPORT_MACHINE, job, "VALIDATED", module="import")
    return job


def preview(session: Session, job_id: int, limit: int = 200, errors_only: bool = False) -> dict:
    job = _job(session, job_id)
    q = select(ImportRow).where(ImportRow.job_id == job_id)
    if errors_only:
        q = q.where(ImportRow.status == "ERROR")
    rows = session.execute(q.order_by(ImportRow.row_no).limit(limit)).scalars().all()
    return {"job": job, "rows": [{"row_no": r.row_no, "status": r.status, "data": json.loads(r.data_json),
                                  "errors": json.loads(r.errors_json) if r.errors_json else []} for r in rows]}


def error_report_xlsx(session: Session, job_id: int) -> bytes:
    rows = session.execute(select(ImportRow).where(ImportRow.job_id == job_id, ImportRow.status == "ERROR")
                           .order_by(ImportRow.row_no)).scalars().all()
    wb = Workbook()
    ws = wb.active
    ws.title = "Errors"
    ws.append(["Excel row", "Errors", "Data"])
    for r in rows:
        ws.append([r.row_no, "; ".join(json.loads(r.errors_json)), r.data_json])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _job(session: Session, job_id: int) -> ImportJob:
    j = session.get(ImportJob, job_id)
    if j is None:
        raise NotFound("Import job not found")
    return j


def submit(session: Session, job: ImportJob, user) -> None:
    if job.created_by_id != user.id:
        raise PermissionDenied("Only the uploader can submit the import")
    if job.error_rows:
        raise BusinessRuleError("The file has validation errors; correct it and upload again", rule_id="IMP-001")
    if job.total_rows == 0:
        raise ValidationFailed("The file contains no data rows")
    job.submitted_by_id = user.id
    sod.record_action(session, "import_job", job.id, "import.submit", user.id)
    transition(session, IMPORT_MACHINE, job, "SUBMITTED", module="import")


def decide(session: Session, job: ImportJob, user, password: str, approve: bool, reason: str) -> None:
    perms = effective_permissions(session, user.id)
    need = ENTITY_APPROVE_PERM.get(job.entity)
    if need and need not in perms:
        raise PermissionDenied(f"Approving a {job.entity} import requires {need}")
    masters.sign_and_transition(session, job, IMPORT_MACHINE, "APPROVED" if approve else "REJECTED", user, password,
                                reason=reason, meaning="APPROVED_BY" if approve else "REJECTED_BY",
                                sod_action="import.approve")
    job.approved_by_id = user.id


def execute(session: Session, job: ImportJob, user) -> dict:
    """Create the records (all-or-nothing). Re-validates first because master data may have changed."""
    if job.status != "APPROVED":
        raise BusinessRuleError("Only an APPROVED import can be executed", rule_id="IMP-002")
    if job.created_by_id != user.id:
        raise PermissionDenied("Only the uploader can execute the approved import")
    rows = [json.loads(r.data_json) for r in session.execute(
        select(ImportRow).where(ImportRow.job_id == job.id).order_by(ImportRow.row_no)).scalars()]
    recheck = validate_rows(session, job.entity, [{k: v for k, v in r.items() if not k.startswith("_")} for r in rows])
    bad = [(i + 2, e) for i, (_c, e) in enumerate(recheck) if e]
    if bad:
        raise BusinessRuleError(f"Data changed since validation; row {bad[0][0]}: {bad[0][1][0]}", rule_id="IMP-003")
    set_reason(f"Controlled import job {job.id}")
    created = _APPLY[job.entity](session, [c for c, _e in recheck])
    for r in session.execute(select(ImportRow).where(ImportRow.job_id == job.id)).scalars():
        r.status = "IMPORTED"
    job.imported_at = utcnow()
    job.result_summary = json.dumps({"created": created})
    transition(session, IMPORT_MACHINE, job, "IMPORTED", module="import")
    return {"created": created}


def _apply_vendor(session, rows):
    n = 0
    for r in rows:
        code = r.pop("vendor_code", None)
        ms.create_vendor(session, {k: v for k, v in r.items() if v is not None}, code)
        n += 1
    return n


def _apply_material(session, rows):
    n = 0
    for r in rows:
        code = r.pop("material_code", None)
        data = {k: v for k, v in r.items() if v is not None and not k.endswith("_code") and not k.startswith("_")}
        data.update(type_id=r["_type_id"], base_unit_id=r["_unit_id"], category_id=r["_category_id"])
        ms.create_material(session, data, code)
        n += 1
    return n


def _apply_location(session, rows):
    n = 0
    for r in rows:
        wh = session.execute(select(Warehouse).where(Warehouse.warehouse_code == r["warehouse_code"])).scalar_one()
        parent = None
        if r.get("parent_code"):
            parent = session.execute(select(Location).where(Location.location_code == r["parent_code"])).scalar_one()
        data = {k: v for k, v in r.items() if v is not None and k not in ("warehouse_code", "parent_code")}
        ms.create_location(session, {**data, "warehouse_id": wh.id, "parent_id": parent.id if parent else None})
        n += 1
    return n


def _apply_stp(session, rows):
    for r in rows:
        key = r.pop("stp_no", None)
        versioning.create_draft(session, STP, {k: v for k, v in r.items() if v is not None}, key)
    return len(rows)


def _apply_spec(session, rows):
    groups: dict[str, list[dict]] = {}
    for r in rows:
        groups.setdefault(r["material_code"], []).append(r)
    for code, items in groups.items():
        m = session.execute(select(Material).where(Material.material_code == code)).scalar_one()
        spec = versioning.create_draft(session, Specification, {"material_id": m.id, "title": f"Specification for {m.name}"})
        for seq, r in enumerate(items, 1):
            session.add(SpecificationParameter(
                specification_id=spec.id, seq=seq, test_name=r["test_name"], test_method=r.get("test_method"),
                stp_id=r.get("_stp_id"), spec_type=r["spec_type"], lsl=r.get("lsl"), usl=r.get("usl"),
                target=r.get("target"), unit=r.get("unit"), decimal_places=r.get("decimal_places"),
                acceptance_criteria=r.get("acceptance_criteria"), criticality=r["criticality"],
                frequency=r.get("frequency"), pharmacopoeial_reference=r.get("pharmacopoeial_reference")))
    return len(groups)


_APPLY = {"vendor": _apply_vendor, "material": _apply_material, "location": _apply_location,
          "stp": _apply_stp, "specification": _apply_spec}
