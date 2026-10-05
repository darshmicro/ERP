"""Certificate of Analysis (spec 43): PDF + Excel, company branding, controlled versions."""
import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.audit import service as audit
from app.core.errors import BusinessRuleError, NotFound
from app.core.time import utcnow
from app.models.audit import ESignature
from app.models.master import Material, Unit, Vendor
from app.models.org import Company
from app.models.platform import Document
from app.models.qc import COA, QCResult, QCTest, Sample
from app.models.spec import STP, SpecificationParameter
from app.models.warehouse import MaterialBatch
from app.services import documents, numbering, qc as qc_svc


def _fmt_limits(r: QCResult) -> str:
    if r.spec_type in ("NUMERIC", "RANGE"):
        lo = f"{float(r.lsl):g}" if r.lsl is not None else None
        hi = f"{float(r.usl):g}" if r.usl is not None else None
        if lo and hi:
            return f"{lo} – {hi} {r.unit or ''}".strip()
        return (f"NLT {lo}" if lo else f"NMT {hi}") + f" {r.unit or ''}".rstrip()
    return r.acceptance_criteria or "Complies"


def _rows(session: Session, lot: MaterialBatch) -> list[dict]:
    rows = []
    for t, r in qc_svc._lot_results(session, lot):
        if r is None:
            continue
        e = qc_svc.effective(session, r)
        stp = session.get(STP, t.stp_id) if t.stp_id else None
        val = e["value"]
        rows.append({"test": t.test_name, "specification": _fmt_limits(r), "result": (f"{float(val):g}" if r.spec_type in ("NUMERIC", "RANGE") and val is not None and not e["amended"]
                                                                                    else str(val)) + (f" {r.unit}" if r.unit and r.spec_type in ("NUMERIC", "RANGE") else ""),
                     "method": f"{stp.stp_no} v{stp.version_no}" if stp else "—", "conforms": "Complies" if e["pass_fail"] == "PASS" else "DOES NOT COMPLY",
                     "amended": e["amended"]})
    return sorted(rows, key=lambda x: x["test"])


def generate(session: Session, lot: MaterialBatch, user, reason: str, signature_id: int | None) -> COA:
    if lot.disposition != "APPROVED":
        raise BusinessRuleError("A CoA can only be issued after QA release (BR-FGR-002)", rule_id="BR-FGR-002")
    rows = _rows(session, lot)
    if not rows:
        raise BusinessRuleError("There are no QC results to certify", rule_id="BR-FGR-002")
    prev = session.execute(select(COA).where(COA.material_batch_id == lot.id).order_by(COA.version_no.desc())).scalars().first()
    coa_no = prev.coa_no if prev else numbering.next_number(session, numbering.default_plant_id(session), "COA")
    ver = (prev.version_no + 1) if prev else 1
    conclusion = "COMPLIES" if all(r["conforms"] == "Complies" for r in rows) else "DOES NOT COMPLY"
    sig = session.get(ESignature, signature_id) if signature_id else None
    ctx = _context(session, lot, rows, coa_no, ver, conclusion, sig)
    pdf, xlsx = render_pdf(ctx), render_xlsx(ctx)
    d_pdf = documents.store(session, f"{coa_no}_v{ver}.pdf", pdf, doc_no=coa_no, version=str(ver))
    d_xl = documents.store(session, f"{coa_no}_v{ver}.xlsx", xlsx, doc_no=coa_no, version=str(ver))
    c = COA(coa_no=coa_no, version_no=ver, material_batch_id=lot.id, pdf_document_id=d_pdf.id, xlsx_document_id=d_xl.id, conclusion=conclusion,
            reason=reason, generated_by_id=user.id if user else None, signature_id=signature_id)
    session.add(c)
    session.flush()
    audit.log_event(session, module="qc", entity="coa", record_id=c.id, action="COA_GENERATED", new=f"{coa_no} v{ver} {conclusion}", reason=reason, signature_id=signature_id)
    return c


def _context(session, lot, rows, coa_no, ver, conclusion, sig) -> dict:
    mat = session.get(Material, lot.material_id)
    comp = session.execute(select(Company)).scalars().first()
    logo = None
    if comp and comp.logo_document_id:
        try:
            logo = documents.read_raw(session.get(Document, comp.logo_document_id))
        except Exception:  # noqa: BLE001
            logo = None
    vendor = session.get(Vendor, lot.vendor_id) if lot.vendor_id else None
    return {"company": comp.name if comp else "GMP-MERP", "address": comp.address if comp else "", "logo": logo, "coa_no": coa_no, "version": ver,
            "product": mat.name, "code": mat.material_code, "lot": lot.lot_no, "vendor": vendor.name if vendor else None, "vendor_batch": lot.vendor_batch_no,
            "mfg": lot.mfg_date, "expiry": lot.expiry_date, "retest": lot.retest_date, "quantity": f"{float(lot.quantity):g} {session.get(Unit, lot.unit_id).code}",
            "rows": rows, "conclusion": conclusion, "qa_release_no": lot.qa_release_no, "qa_date": lot.released_at,
            "approved_by": sig.printed_name if sig else None, "manifest": sig.manifest_hash if sig else None, "role": sig.role_name if sig else None,
            "qc_no": lot.qc_no, "generated": utcnow()}


def _d(x) -> str:
    return x.strftime("%d-%b-%Y") if x else "—"


def render_pdf(c: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=18 * mm, bottomMargin=20 * mm,
                            title=f"CoA {c['coa_no']} v{c['version']}")
    st = getSampleStyleSheet()
    els = []
    head = [[Image(io.BytesIO(c["logo"]), 22 * mm, 22 * mm, kind="proportional") if c["logo"] else "", Paragraph(f"<b>{c['company']}</b><br/>{c['address'] or ''}", st["Normal"]),
             Paragraph(f"<b>CERTIFICATE OF ANALYSIS</b><br/>No. {c['coa_no']}<br/>Version {c['version']}", st["Normal"])]]
    t = Table(head, colWidths=[28 * mm, 90 * mm, 60 * mm])
    t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEBELOW", (0, 0), (-1, 0), 1, colors.black)]))
    els += [t, Spacer(1, 4 * mm)]
    info = [["Product", c["product"], "Product code", c["code"]], ["Batch / lot", c["lot"], "QC no.", c["qc_no"] or "—"],
            ["Vendor", c["vendor"] or "—", "Vendor batch", c["vendor_batch"] or "—"], ["Mfg date", _d(c["mfg"]), "Expiry date", _d(c["expiry"])],
            ["Retest date", _d(c["retest"]), "Quantity", c["quantity"]]]
    it = Table(info, colWidths=[28 * mm, 62 * mm, 28 * mm, 60 * mm])
    it.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 8.5), ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"), ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
                            ("GRID", (0, 0), (-1, -1), 0.3, colors.grey)]))
    els += [it, Spacer(1, 4 * mm)]
    data = [["Test", "Specification", "Result", "Method", "Conclusion"]] + [[r["test"], r["specification"], r["result"] + (" *" if r["amended"] else ""), r["method"], r["conforms"]] for r in c["rows"]]
    tt = Table(data, colWidths=[40 * mm, 45 * mm, 33 * mm, 30 * mm, 30 * mm], repeatRows=1)
    tt.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f3a5f")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTSIZE", (0, 0), (-1, -1), 8),
                            ("GRID", (0, 0), (-1, -1), 0.3, colors.grey), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    els += [tt, Spacer(1, 4 * mm), Paragraph(f"<b>Conclusion: {c['conclusion']}</b>", st["Normal"]), Spacer(1, 2 * mm)]
    if any(r["amended"] for r in c["rows"]):
        els.append(Paragraph("* result corrected through the controlled amendment procedure", st["Italic"]))
    sig = [[Paragraph(f"<b>QA release:</b> {c['qa_release_no'] or '—'}<br/><b>Released on:</b> {_d(c['qa_date'])}<br/><b>Approved by (e-signature):</b> {c['approved_by'] or '—'}"
                      f"<br/><b>Role:</b> {c['role'] or '—'}<br/><span size=6>Signature manifest {c['manifest'] or '—'}</span>", st["Normal"])]]
    els += [Spacer(1, 6 * mm), Table(sig, colWidths=[178 * mm], style=[("BOX", (0, 0), (-1, -1), 0.5, colors.black)])]

    def footer(cv: canvas.Canvas, d):
        cv.setFont("Helvetica", 7)
        cv.drawString(15 * mm, 10 * mm, f"CoA {c['coa_no']} v{c['version']} · generated {c['generated'].strftime('%d-%b-%Y %H:%M')} UTC · Page {d.page}")
        cv.drawRightString(195 * mm, 10 * mm, "Electronically signed record — valid without handwritten signature")
    doc.build(els, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


def render_xlsx(c: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "CoA"
    ws.append([c["company"]])
    ws.append([f"CERTIFICATE OF ANALYSIS {c['coa_no']} v{c['version']}"])
    ws["A1"].font, ws["A2"].font = Font(bold=True, size=14), Font(bold=True, size=12)
    for k, v in (("Product", c["product"]), ("Product code", c["code"]), ("Batch / lot", c["lot"]), ("Vendor batch", c["vendor_batch"] or "—"),
                 ("Mfg date", _d(c["mfg"])), ("Expiry date", _d(c["expiry"])), ("QC no.", c["qc_no"] or "—")):
        ws.append([k, v])
    ws.append([])
    ws.append(["Test", "Specification", "Result", "Method", "Conclusion"])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    for r in c["rows"]:
        ws.append([r["test"], r["specification"], r["result"], r["method"], r["conforms"]])
    ws.append([])
    ws.append(["Conclusion", c["conclusion"]])
    ws.append(["QA release no.", c["qa_release_no"] or "—"])
    ws.append(["Released on", _d(c["qa_date"])])
    ws.append(["Approved by (e-signature)", c["approved_by"] or "—"])
    ws.append(["Signature manifest", c["manifest"] or "—"])
    for col, w in zip("ABCDE", (30, 34, 24, 18, 20)):
        ws.column_dimensions[col].width = w
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def latest(session: Session, lot_id: int) -> COA | None:
    return session.execute(select(COA).where(COA.material_batch_id == lot_id).order_by(COA.version_no.desc())).scalars().first()
