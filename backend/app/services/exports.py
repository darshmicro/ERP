"""Excel export with report header (spec 66): title, generated date/time, filters, user, then data."""
import io
from datetime import datetime, timezone
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def build_xlsx(title: str, filters: dict[str, Any], columns: Sequence[tuple[str, str]], rows: Sequence[dict],
               generated_by: str, company: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([company or "GMP-MERP"])
    ws.append([title])
    ws.append([f"Generated: {datetime.now(timezone.utc).strftime('%d-%b-%Y %H:%M:%S')} UTC   By: {generated_by}"])
    ws.append(["Filters: " + (", ".join(f"{k}={v}" for k, v in filters.items() if v not in (None, "")) or "none")])
    ws.append([])
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"].font = Font(bold=True, size=12)
    header_row = ws.max_row + 1
    ws.append([label for _k, label in columns])
    for c in ws[header_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in rows:
        ws.append([_cell(r.get(k)) for k, _l in columns])
    for i, (k, label) in enumerate(columns, 1):
        width = max([len(label)] + [len(str(_cell(r.get(k)) or "")) for r in rows[:200]]) + 2
        ws.column_dimensions[ws.cell(row=header_row, column=i).column_letter].width = min(width, 50)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.append([])
    ws.append([f"Records: {len(rows)}"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell(v: Any) -> Any:
    if isinstance(v, datetime) and v.tzinfo:
        return v.astimezone(timezone.utc).replace(tzinfo=None)
    if isinstance(v, bool):
        return "Yes" if v else "No"
    if hasattr(v, "__float__") and not isinstance(v, (int, float)):
        return float(v)
    return v
