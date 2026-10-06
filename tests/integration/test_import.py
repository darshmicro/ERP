import io

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.core import db
from app.models import Location, Material, Specification, STP, Vendor
from tests.conftest import PW
from tests.helpers import as_user, bootstrap_basics, xlsx

XL = ("file", ("f.xlsx", b"", "application/octet-stream"))


def _upload(c, h, entity, content, name="f.xlsx"):
    return c.post("/api/v1/imports", headers=h, data={"entity": entity}, files={"file": (name, content, "application/octet-stream")})


@pytest.fixture()
def w(app):
    ids = bootstrap_basics(app)
    pu, hp = as_user(app, "imp_pu", ["PURCHASE_USER"])
    qa, hq = as_user(app, "imp_qa", ["QA_HEAD"])
    return dict(ids=ids, pu=pu, hp=hp, qa=qa, hq=hq, app=app)


def test_template_download_has_columns_and_instructions(w):
    r = w["pu"].get("/api/v1/imports/templates/vendor")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb["Data"][1][0].value == "name" and "Instructions" in wb.sheetnames
    assert w["pu"].get("/api/v1/imports/templates/bom").status_code == 422   # BOM import arrives with Phase 6


def test_full_import_flow_vendor(w):
    data = xlsx(["name", "risk_class", "email", "city"], [["Vendor A", "high", "a@x.com", "Pune"], ["Vendor B", None, None, "Mumbai"]])
    r = _upload(w["pu"], w["hp"], "vendor", data)
    assert r.status_code == 201, r.text
    job = r.json()
    assert job["status"] == "VALIDATED" and job["total_rows"] == 2 and job["error_rows"] == 0
    s = db.new_session()
    assert s.execute(select(func.count()).select_from(Vendor)).scalar() == 0          # nothing in production tables yet
    s.close()
    prev = w["pu"].get(f"/api/v1/imports/{job['id']}").json()
    assert prev["rows"][0]["data"]["risk_class"] == "HIGH"
    # cannot execute before approval, cannot approve before submission
    assert w["pu"].post(f"/api/v1/imports/{job['id']}/execute", headers=w["hp"]).status_code == 409
    assert w["qa"].post(f"/api/v1/imports/{job['id']}/decision", headers=w["hq"], json={"password": PW, "reason": "ok"}).status_code == 409
    assert w["pu"].post(f"/api/v1/imports/{job['id']}/submit", headers=w["hp"]).status_code == 200
    # uploader has no approval permission; wrong password refused; QA approves with e-signature
    assert w["pu"].post(f"/api/v1/imports/{job['id']}/decision", headers=w["hp"], json={"password": PW, "reason": "self"}).status_code == 403
    assert w["qa"].post(f"/api/v1/imports/{job['id']}/decision", headers=w["hq"], json={"password": "bad-Password-1!", "reason": "ok"}).status_code == 401
    assert w["qa"].post(f"/api/v1/imports/{job['id']}/decision", headers=w["hq"], json={"password": PW, "reason": "data verified"}).status_code == 200
    # only the uploader executes
    assert w["qa"].post(f"/api/v1/imports/{job['id']}/execute", headers=w["hq"]).status_code == 403
    ex = w["pu"].post(f"/api/v1/imports/{job['id']}/execute", headers=w["hp"])
    assert ex.status_code == 200 and ex.json()["created"] == 2
    s = db.new_session()
    vs = s.execute(select(Vendor).order_by(Vendor.id)).scalars().all()
    assert [v.approval_status for v in vs] == ["DRAFT", "DRAFT"] and vs[0].vendor_code == "VEN-00001"
    s.close()
    assert w["pu"].post(f"/api/v1/imports/{job['id']}/execute", headers=w["hp"]).status_code == 409   # not repeatable


def test_validation_errors_block_submission_and_report(w):
    data = xlsx(["name", "risk_class", "email", "pan_no"],
                [["Good Vendor", "LOW", "g@x.com", None], [None, "LOW", None, None], ["Bad Risk", "EXTREME", "nope", "xx"],
                 ["Good Vendor", None, None, None]])
    job = _upload(w["pu"], w["hp"], "vendor", data).json()
    assert job["error_rows"] == 3
    rows = w["pu"].get(f"/api/v1/imports/{job['id']}", params={"errors_only": True}).json()["rows"]
    assert {r["row_no"] for r in rows} == {3, 4, 5}
    assert any("required" in e for e in rows[0]["errors"])
    rep = w["pu"].get(f"/api/v1/imports/{job['id']}/error-report")
    ws = load_workbook(io.BytesIO(rep.content)).active
    assert ws.max_row == 4
    r = w["pu"].post(f"/api/v1/imports/{job['id']}/submit", headers=w["hp"])
    assert r.status_code == 409 and r.json()["rule_id"] == "IMP-001"


def test_upload_rejects_wrong_format_and_missing_columns(w):
    assert _upload(w["pu"], w["hp"], "vendor", b"name\nx", "x.csv").status_code == 422
    assert _upload(w["pu"], w["hp"], "vendor", b"not an excel", "x.xlsx").status_code == 422
    assert _upload(w["pu"], w["hp"], "vendor", xlsx(["city"], [["Pune"]])).status_code == 422
    assert _upload(w["pu"], w["hp"], "material", xlsx(["name"], [["x"]])).status_code == 422


def _run(w, entity, data, creator=("pu", "hp")):
    c, h = w[creator[0]], w[creator[1]]
    job = _upload(c, h, entity, data).json()
    assert job["error_rows"] == 0, c.get(f"/api/v1/imports/{job['id']}").json()
    c.post(f"/api/v1/imports/{job['id']}/submit", headers=h)
    assert w["qa"].post(f"/api/v1/imports/{job['id']}/decision", headers=w["hq"], json={"password": PW, "reason": "ok"}).status_code == 200
    return c.post(f"/api/v1/imports/{job['id']}/execute", headers=h)


def test_material_location_stp_and_spec_imports(w):
    qc, hqc = as_user(w["app"], "imp_qc", ["QC_ANALYST"])
    w["qc"], w["hqc"] = qc, hqc
    bad = _upload(qc, hqc, "material", xlsx(["name", "type_code", "base_unit_code"], [["Mat X", "ZZ", "kg"]])).json()
    assert bad["error_rows"] == 1
    r = _run(w, "material", xlsx(["name", "type_code", "base_unit_code", "category_code", "temp_min", "temp_max", "shelf_life_days"],
                                 [["Sodium Chloride", "RM", "kg", "CHEM", 15, 25, 730], ["Glycine", "rm", "KG", None, None, None, None]]),
             creator=("qc", "hqc"))
    assert r.status_code == 200 and r.json()["created"] == 2
    s = db.new_session()
    mats = s.execute(select(Material).order_by(Material.id)).scalars().all()
    assert [m.master_status for m in mats] == ["DRAFT", "DRAFT"] and mats[0].material_code == "MAT-00001"
    s.close()
    # spec import needs APPROVED/ACTIVE material -> error for DRAFT
    e = _upload(qc, hqc, "specification", xlsx(["material_code", "test_name", "lsl", "usl"], [["MAT-00001", "Assay", 98, 101]])).json()
    assert e["error_rows"] == 1
    # stp import (draft), then location import with parents
    r = _run(w, "stp", xlsx(["title", "procedure"], [["Assay STP", "Do the assay"]]), creator=("qc", "hqc"))
    assert r.status_code == 200
    wh, hw = as_user(w["app"], "imp_wh", ["WAREHOUSE_USER"])
    wh.post("/api/v1/warehouses", headers=hw, json={"warehouse_code": "WH9", "name": "Store", "reason": "x"})
    w["wh"], w["hw"] = wh, hw
    locs = xlsx(["warehouse_code", "location_code", "name", "location_type", "parent_code"],
                [["WH9", "Z9", "Zone 9", "ZONE", None], ["WH9", "RM9", "Room 9", "ROOM", "Z9"], ["WH9", "RK9", "Rack 9", "RACK", "RM9"]])
    # location import requires import.job.create: warehouse user lacks it -> forbidden
    assert _upload(wh, hw, "location", locs).status_code == 403
    pu_loc = _upload(w["pu"], w["hp"], "location", xlsx(["warehouse_code", "location_code", "name", "location_type", "parent_code"],
                                                          [["WH9", "RK9", "Rack", "RACK", "RM9"]])).json()
    assert pu_loc["error_rows"] == 1                                    # parent defined later / missing
    r = _run(w, "location", locs)
    assert r.status_code == 200 and r.json()["created"] == 3
    s = db.new_session()
    assert s.execute(select(func.count()).select_from(Location)).scalar() == 3
    s.close()


def test_import_rows_cannot_be_executed_if_data_changed_since_validation(w):
    data = xlsx(["name"], [["Race Vendor"]])
    job = _upload(w["pu"], w["hp"], "vendor", data).json()
    w["pu"].post(f"/api/v1/imports/{job['id']}/submit", headers=w["hp"])
    w["qa"].post(f"/api/v1/imports/{job['id']}/decision", headers=w["hq"], json={"password": PW, "reason": "ok"})
    w["pu"].post("/api/v1/vendors", headers=w["hp"], json={"name": "Race Vendor", "reason": "created manually meanwhile"})
    r = w["pu"].post(f"/api/v1/imports/{job['id']}/execute", headers=w["hp"])
    assert r.status_code == 409 and r.json()["rule_id"] == "IMP-003"
