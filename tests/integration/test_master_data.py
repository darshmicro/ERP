import pytest
from sqlalchemy import select

from app.core import db
from app.models import AuditTrail, ESignature
from tests.conftest import PW
from tests.helpers import PDF, as_user, bootstrap_basics


@pytest.fixture()
def ids(app):
    return bootstrap_basics(app)


def test_admin_manages_lookups_with_reason_and_others_cannot(app, ids):
    pu, hp = as_user(app, "pu1", ["PURCHASE_USER"])
    assert pu.post("/api/v1/units", headers=hp, json={"code": "g", "name": "Gram", "reason": "x"}).status_code == 403
    a, ha = as_user(app, "admin_b", [])
    assert a.patch(f"/api/v1/units/{ids['kg']}", headers=ha, json={"name": "Kilo"}).status_code == 422  # no reason
    r = a.patch(f"/api/v1/units/{ids['kg']}", headers=ha, json={"name": "Kilo", "reason": "rename"})
    assert r.status_code == 200 and r.json()["name"] == "Kilo"
    assert a.post("/api/v1/units", headers=ha, json={"code": "KG", "name": "dup", "reason": "x"}).status_code == 409
    assert a.post("/api/v1/units", headers=ha, json={"code": "x", "name": "x", "hacked": 1, "reason": "x"}).status_code == 422


def test_vendor_lifecycle_sod_and_signature(app, ids):
    pm, hpm = as_user(app, "pm1", ["PURCHASE_MANAGER"])
    r = pm.post("/api/v1/vendors", headers=hpm, json={"name": "Vendor A", "risk_class": "HIGH", "bank_account_no": "123456789012",
                                                     "country": "India", "reason": "new supplier"})
    assert r.status_code == 201, r.text
    v = r.json()
    assert v["vendor_code"] == "VEN-00001" and v["approval_status"] == "DRAFT"
    assert v["bank_account_masked"].endswith("9012") and "bank_account_no" not in v
    assert pm.post("/api/v1/vendors", headers=hpm, json={"name": "vendor a", "reason": "dup"}).status_code == 409
    assert pm.post("/api/v1/vendors", headers=hpm, json={"name": "Vendor B", "pan_no": "bad", "reason": "x"}).status_code == 422
    # author cannot approve own record (SOD-13)
    r = pm.post(f"/api/v1/vendors/{v['id']}/approve", headers=hpm, json={"password": PW, "reason": "ok"})
    assert r.status_code == 409 and r.json()["rule_id"] == "SOD-13"
    pm2, h2 = as_user(app, "pm2", ["PURCHASE_MANAGER"])
    bad = pm2.post(f"/api/v1/vendors/{v['id']}/approve", headers=h2, json={"password": "wrong-Password-1!", "reason": "ok"})
    assert bad.status_code == 401
    ok = pm2.post(f"/api/v1/vendors/{v['id']}/approve", headers=h2, json={"password": PW, "reason": "documents reviewed"})
    assert ok.status_code == 200 and ok.json()["approval_status"] == "APPROVED"
    s = db.new_session()
    sig = s.execute(select(ESignature).where(ESignature.entity == "vendor")).scalars().one()
    assert sig.meaning == "APPROVED_BY" and sig.username == "pm2"
    audit_vals = " ".join(str(a.new_value) for a in s.execute(select(AuditTrail).where(AuditTrail.entity == "vendor")).scalars())
    assert "123456789012" not in audit_vals
    s.close()
    # status is not writable through PATCH
    assert pm2.patch(f"/api/v1/vendors/{v['id']}", headers=h2, json={"approval_status": "DRAFT", "reason": "x"}).status_code == 422


def test_vendor_documents_review_and_integrity(app, ids):
    qa, hq = as_user(app, "qa_o", ["QA_OFFICER"])
    v = qa.post("/api/v1/vendors", headers=hq, json={"name": "Doc Vendor", "reason": "x"}).json()
    up = lambda **kw: qa.post(f"/api/v1/vendors/{v['id']}/documents", headers=hq, data=kw, files={"file": ("gmp.pdf", PDF, "application/pdf")})
    assert up(doc_type="NOPE").status_code == 422
    bad = qa.post(f"/api/v1/vendors/{v['id']}/documents", headers=hq, data={"doc_type": "GMP_CERTIFICATE"},
                  files={"file": ("gmp.pdf", b"not a pdf", "application/pdf")})
    assert bad.status_code == 422
    r = up(doc_type="GMP_CERTIFICATE", version="1", expiry_date="2020-01-01", issue_date="2019-01-01")
    assert r.status_code == 201, r.text
    docs = qa.get(f"/api/v1/vendors/{v['id']}").json()["documents"]
    assert docs[0]["expired"] is True and docs[0]["review_status"] == "PENDING"
    vd = docs[0]["id"]
    # uploader cannot review own upload (SOD-19)
    r = qa.post(f"/api/v1/vendor-documents/{vd}/review", headers=hq, json={"approve": True, "comment": "ok"})
    assert r.status_code == 409 and r.json()["rule_id"] == "SOD-19"
    qh, hh = as_user(app, "qa_h2", ["QA_OFFICER"])
    assert qh.post(f"/api/v1/vendor-documents/{vd}/review", headers=hh, json={"approve": True, "comment": "valid"}).status_code == 200
    # download is hash-verified and audited
    d = qa.get(f"/api/v1/documents/{docs[0]['document_id']}/download")
    assert d.status_code == 200 and d.content == PDF
    s = db.new_session()
    assert s.execute(select(AuditTrail).where(AuditTrail.action == "DOWNLOAD")).first()
    # tamper with stored file -> integrity failure
    import os
    from app.core.config import get_settings
    from app.models import Document
    doc = s.get(Document, docs[0]["document_id"])
    with open(os.path.join(get_settings().file_storage_path, doc.storage_key), "wb") as fh:
        fh.write(PDF + b"tampered")
    s.close()
    assert qa.get(f"/api/v1/documents/{docs[0]['document_id']}/download").json()["code"] == "DOC_INTEGRITY"


def _material(c, h, ids, **kw):
    body = {"name": "Sodium Chloride IP", "type_id": ids["RM"], "base_unit_id": ids["kg"], "category_id": ids["chem"],
            "reason": "new material", **kw}
    return c.post("/api/v1/materials", headers=h, json=body)


def test_material_validation_and_lifecycle(app, ids):
    qc, hq = as_user(app, "qc_a", ["QC_ANALYST"])
    r = _material(qc, hq, ids, temp_min=25, temp_max=2)
    assert r.status_code == 422
    assert _material(qc, hq, ids, shelf_life_days=100, retest_days=200).status_code == 422
    assert _material(qc, hq, ids, type_id=9999).status_code == 422  # unknown FK -> friendly error
    r = _material(qc, hq, ids, shelf_life_days=730, retest_days=365, storage_condition="15-25C")
    assert r.status_code == 201, r.text
    m = r.json()
    assert m["material_code"] == "MAT-00001" and m["barcode"] == "MAT-00001" and m["master_status"] == "DRAFT"
    # the QC analyst cannot approve
    assert qc.post(f"/api/v1/materials/{m['id']}/approve", headers=hq, json={"password": PW, "reason": "x"}).status_code == 403
    qa, hqa = as_user(app, "qa_head1", ["QA_HEAD"])
    r = qa.post(f"/api/v1/materials/{m['id']}/approve", headers=hqa, json={"password": PW, "reason": "specs reviewed"})
    assert r.status_code == 200 and r.json()["master_status"] == "APPROVED"
    # cannot skip: DRAFT -> OBSOLETE style jumps are not endpoints; approved -> approve again is illegal
    assert qa.post(f"/api/v1/materials/{m['id']}/approve", headers=hqa, json={"password": PW, "reason": "again"}).status_code == 409
    assert qa.post(f"/api/v1/materials/{m['id']}/activate", headers=hqa, json={"password": PW, "reason": "go live"}).json()["master_status"] == "ACTIVE"
    # edits need a reason and are audited with old/new
    assert qc.patch(f"/api/v1/materials/{m['id']}", headers=hq, json={"grade": "IP"}).status_code == 422
    assert qc.patch(f"/api/v1/materials/{m['id']}", headers=hq, json={"grade": "IP", "reason": "grade added"}).status_code == 200
    s = db.new_session()
    row = s.execute(select(AuditTrail).where(AuditTrail.entity == "material", AuditTrail.field_name == "grade")).scalars().one()
    assert (row.old_value, row.new_value, row.reason) == (None, "IP", "grade added")
    s.close()
    # status and code cannot be patched directly
    assert qc.patch(f"/api/v1/materials/{m['id']}", headers=hq, json={"master_status": "DRAFT", "reason": "x"}).status_code == 422
    assert qc.patch(f"/api/v1/materials/{m['id']}", headers=hq, json={"material_code": "HACK", "reason": "x"}).status_code == 422
    # obsolete: frozen
    qa.post(f"/api/v1/materials/{m['id']}/obsolete", headers=hqa, json={"password": PW, "reason": "replaced"})
    assert qc.patch(f"/api/v1/materials/{m['id']}", headers=hq, json={"grade": "X", "reason": "late"}).status_code == 403


def test_material_author_cannot_approve_sod14(app, ids):
    qa, h = as_user(app, "qa_head2", ["QA_HEAD"])
    m = _material(qa, h, ids, name="Glycine").json()
    r = qa.post(f"/api/v1/materials/{m['id']}/approve", headers=h, json={"password": PW, "reason": "self"})
    assert r.status_code == 409 and r.json()["rule_id"] == "SOD-14"


def test_role_access_matrix_phase2(app, ids):
    wh, hw = as_user(app, "wh_x", ["WAREHOUSE_USER"])
    assert wh.post("/api/v1/vendors", headers=hw, json={"name": "X", "reason": "r"}).status_code == 403
    assert wh.get("/api/v1/materials").status_code == 200
    pr, hp = as_user(app, "prod_x", ["PRODUCTION_USER"])
    assert pr.post("/api/v1/materials", headers=hp, json={"name": "Y", "type_id": 1, "base_unit_id": 1, "reason": "r"}).status_code == 403
    ds, hd = as_user(app, "dsp_x", ["DISPATCH_USER"])
    assert ds.post("/api/v1/customers", headers=hd, json={"name": "Customer One", "reason": "r"}).json()["customer_code"] == "CUS-00001"
    ad, ha = as_user(app, "admin_b", [])
    assert ad.get("/api/v1/vendors").status_code == 403          # admin has no business-data access
    assert ad.get("/api/v1/materials").status_code == 403
    aud, hu = as_user(app, "auditor_x", ["AUDITOR"])
    assert aud.get("/api/v1/vendors").status_code == 200
    assert aud.post("/api/v1/vendors", headers=hu, json={"name": "Z", "reason": "r"}).status_code == 403


def test_exports_have_report_header_and_are_audited(app, ids):
    import io
    from openpyxl import load_workbook
    qa, h = as_user(app, "qa_exp", ["QA_OFFICER"])
    qa.post("/api/v1/vendors", headers=h, json={"name": "Export Vendor", "bank_account_no": "999988887777", "reason": "x"})
    r = qa.get("/api/v1/vendors/export", params={"q": "export"})
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    cells = [c for row in ws.iter_rows(values_only=True) for c in row if c]
    text = " ".join(str(c) for c in cells)
    assert "Vendor master" in text and "Generated:" in text and "qa_exp" in text and "Filters:" in text
    assert "Export Vendor" in text and "999988887777" not in text   # unauthorised/sensitive fields are not exported
    s = db.new_session()
    assert s.execute(select(AuditTrail).where(AuditTrail.action == "EXPORT", AuditTrail.entity == "vendor")).first()
    s.close()
    wh, hw = as_user(app, "wh_exp", ["PURCHASE_USER"])
    assert wh.get("/api/v1/vendors/export").status_code == 200  # purchase has export perm via MD_READ


def test_seed_upgrade_only_grants_new_permissions_and_never_undoes_admin_changes(app, ids):
    from app.audit.context import AuditContext, audit_context
    from app.models import Permission, Role, RolePermission
    from app.services import seed
    s = db.new_session()
    role = s.execute(select(Role).where(Role.role_code == "PURCHASE_USER")).scalar_one()
    perm = s.execute(select(Permission).where(Permission.perm_code == "md.vendor.update")).scalar_one()
    with audit_context(AuditContext(user_name="t", reason="admin revoked")):
        from app.core.time import utcnow
        rp = s.execute(select(RolePermission).where(RolePermission.role_id == role.id, RolePermission.permission_id == perm.id)).scalar_one()
        rp.revoked_at = utcnow()
        s.commit()
        seed.seed_baseline(s)       # re-running the seed must not re-grant a deliberately revoked permission
        s.commit()
    s.refresh(rp)
    assert rp.revoked_at is not None
    s.close()
