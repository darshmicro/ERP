import threading
from datetime import date, timedelta
from decimal import Decimal

import pytest
from sqlalchemy import select, text

from app.core import db
from app.core.errors import BusinessRuleError, ImmutableRecordError
from app.models import (AuditTrail, InventoryBalance, InventoryTransaction, MaterialBatch, MaterialLabel, Notification,
                        PurchaseOrder, QualityHold)
from app.services import inventory, lots
from tests.conftest import PW
from tests.helpers import (as_user, answer_all, build_warehouse_world, grn_body, make_lot, raw, received_lot)

GRN = "/api/v1/grn"


@pytest.fixture()
def w(app):
    return build_warehouse_world(app)


def _lot(lid):
    s = db.new_session()
    lot = s.get(MaterialBatch, lid)
    s.close()
    return lot


# ------------------------------------------------------------------ GRN creation rules
def test_grn_requires_approved_po_and_blocks_over_delivery_and_bad_dates(w):
    c, h = w["wh1"], w["hwh1"]
    r = c.post(GRN, headers=h, json=grn_body(w, qty=101))
    assert r.status_code == 409 and r.json()["rule_id"] == "BR-GRN-001"                       # over-delivery
    exp = c.post(GRN, headers=h, json=grn_body(w, line={"expiry_date": str(date.today() - timedelta(days=1))}))
    assert exp.status_code == 409 and exp.json()["rule_id"] == "BR-GRN-003"                    # already expired
    rev = c.post(GRN, headers=h, json=grn_body(w, line={"mfg_date": str(date.today()), "expiry_date": str(date.today() - timedelta(days=5))}))
    assert rev.status_code == 409
    assert c.post(GRN, headers=h, json=grn_body(w, line={"vendor_batch_no": " "})).status_code in (409, 422)
    # a draft PO cannot be received against
    po2 = w["pu"].post("/api/v1/purchase-orders", headers=w["hpu"], json={"vendor_id": w["vendor"], "reason": "x", "lines": [{"material_id": w["material"], "quantity": 5, "rate": 1}]}).json()
    body = grn_body(w)
    body["po_id"] = po2["id"]
    body["lines"][0]["po_line_id"] = po2["lines"][0]["id"]
    assert c.post(GRN, headers=h, json=body).status_code == 409


def test_receipt_creates_quarantine_lot_containers_ledger_and_closes_po(w):
    out = received_lot(w)
    lot = out["lots"][0]
    assert out["status"] == "QUARANTINE" and lot["disposition"] == "QUARANTINE" and lot["lot_no"].startswith("LOT-")
    d = w["wh1"].get(f"/api/v1/lots/{lot['id']}").json()
    assert d["status"] == "QUARANTINE" and d["on_hand"] == 100.0 and d["balances"][0]["location"] == "QZ1"
    assert len(d["containers"]) == 4 and sum(c["quantity"] for c in d["containers"]) == 100.0
    assert d["grn"]["po_no"].startswith("PO-") and d["specification_id"] == w["spec"]
    led = w["wh1"].get(f"/api/v1/lots/{lot['id']}/ledger").json()
    assert [t["txn_type"] for t in led] == ["RECEIPT"] and led[0]["signature_id"]
    po = w["pu"].get(f"/api/v1/purchase-orders/{w['po']}").json()
    assert po["status"] == "CLOSED" and float(po["lines"][0]["received_quantity"]) == 100.0
    assert w["wh1"].get("/api/v1/dashboard/summary").json()["cards"]["quarantine_materials"]["value"] == 1


def test_partial_receipts_and_po_status(w):
    received_lot(w, qty=60)
    assert w["pu"].get(f"/api/v1/purchase-orders/{w['po']}").json()["status"] == "PARTIALLY_RECEIVED"
    assert w["wh1"].post(GRN, headers=w["hwh1"], json=grn_body(w, qty=50)).status_code == 409        # only 40 left
    received_lot(w, qty=40)
    assert w["pu"].get(f"/api/v1/purchase-orders/{w['po']}").json()["status"] == "CLOSED"


# ------------------------------------------------------------------ verification & checklist
def test_second_person_verification_and_incomplete_checklist(w):
    g = w["wh1"].post(GRN, headers=w["hwh1"], json=grn_body(w)).json()
    assert w["wh2"].post(f"{GRN}/{g['id']}/submit", headers=w["hwh2"]).status_code == 403          # only the receiver submits
    w["wh1"].post(f"{GRN}/{g['id']}/submit", headers=w["hwh1"])
    verify = lambda c, h: c.post(f"{GRN}/{g['id']}/verify", headers=h, json={"quarantine_location_id": w["qloc"], "password": PW, "reason": "ok"})
    r = verify(w["wh2"], w["hwh2"])
    assert r.status_code == 409 and r.json()["rule_id"] == "BR-GRN-004" and "incomplete" in r.json()["message"].lower()
    answer_all(w, w["wh1"], w["hwh1"], g["id"])
    same = verify(w["wh1"], w["hwh1"])
    assert same.status_code == 409 and same.json()["rule_id"] == "SOD-22"                           # receiver cannot verify
    assert verify(w["wh2"], w["hwh2"]).status_code == 200


def test_critical_checklist_failure_blocks_receipt_until_qa_exception(w):
    g = w["wh1"].post(GRN, headers=w["hwh1"], json=grn_body(w)).json()
    w["wh1"].post(f"{GRN}/{g['id']}/submit", headers=w["hwh1"])
    st = answer_all(w, w["wh2"], w["hwh2"], g["id"], overrides={"COA_AVAILABLE": "NO"})
    assert st["critical_failures"] == ["CoA available"]
    ver = lambda: w["wh2"].post(f"{GRN}/{g['id']}/verify", headers=w["hwh2"], json={"quarantine_location_id": w["qloc"], "password": PW, "reason": "ok"})
    r = ver()
    assert r.status_code == 409 and "CoA available" in r.json()["message"]
    item = next(i for i in st["items"] if i["code"] == "COA_AVAILABLE")["item_id"]
    # warehouse cannot grant exceptions; QA head can, with e-signature + reference
    assert w["wh2"].post(f"{GRN}/{g['id']}/exception", headers=w["hwh2"], json={"item_id": item, "exception_ref": "DEV-1", "password": PW, "reason": "x"}).status_code == 403
    assert w["qa"].post(f"{GRN}/{g['id']}/exception", headers=w["hqa"], json={"item_id": item, "exception_ref": "DEV-1", "password": "bad-Password-1!", "reason": "x"}).status_code == 401
    ok = w["qa"].post(f"{GRN}/{g['id']}/exception", headers=w["hqa"], json={"item_id": item, "exception_ref": "DEV-2026-001", "password": PW, "reason": "CoA emailed by vendor, verified"})
    assert ok.status_code == 200 and ok.json()["critical_failures"] == []
    assert ver().status_code == 200


def test_qa_can_reject_receipt(w):
    g = w["wh1"].post(GRN, headers=w["hwh1"], json=grn_body(w)).json()
    w["wh1"].post(f"{GRN}/{g['id']}/submit", headers=w["hwh1"])
    assert w["wh2"].post(f"{GRN}/{g['id']}/reject", headers=w["hwh2"], json={"password": PW, "reason": "damaged"}).status_code == 403
    r = w["qa"].post(f"{GRN}/{g['id']}/reject", headers=w["hqa"], json={"password": PW, "reason": "container damaged"})
    assert r.status_code == 200 and r.json()["status"] == "REJECTED"
    s = db.new_session()
    assert s.execute(select(MaterialBatch)).first() is None          # nothing entered stock
    s.close()


def test_quarantine_location_required_and_vendor_expiry_at_receipt_places_hold(w):
    g = w["wh1"].post(GRN, headers=w["hwh1"], json=grn_body(w)).json()
    w["wh1"].post(f"{GRN}/{g['id']}/submit", headers=w["hwh1"])
    answer_all(w, w["wh2"], w["hwh2"], g["id"])
    r = w["wh2"].post(f"{GRN}/{g['id']}/verify", headers=w["hwh2"], json={"quarantine_location_id": w["aloc"], "password": PW, "reason": "ok"})
    assert r.status_code == 409 and r.json()["rule_id"] == "BR-QRN-001"
    raw("UPDATE vendor_qualification SET requalification_due_date=:d", d=str(date.today() - timedelta(days=1)))   # vendor lapsed in transit
    r = w["wh2"].post(f"{GRN}/{g['id']}/verify", headers=w["hwh2"], json={"quarantine_location_id": w["qloc"], "password": PW, "reason": "ok"})
    assert r.status_code == 200
    lot = r.json()["lots"][0]
    d = w["wh1"].get(f"/api/v1/lots/{lot['id']}").json()
    assert d["status"] == "HOLD" and d["holds"][0]["source"] == "GRN"


# ------------------------------------------------------------------ crit tests 4, 5, 7 and holds
def test_crit_04_quarantine_material_cannot_be_issued(w):
    lot = received_lot(w)["lots"][0]
    d = w["wh1"].get(f"/api/v1/lots/{lot['id']}").json()
    assert d["issue_violations"][0]["rule_id"] == "BR-ISS-001"
    s = db.new_session()
    with pytest.raises(BusinessRuleError) as e:
        lots.assert_issuable(s, s.get(MaterialBatch, lot["id"]))
    assert e.value.rule_id == "BR-ISS-001"
    s.close()


def test_crit_05_rejected_and_crit_07_expired_material_cannot_be_issued(w):
    ids = w["ids"]
    rej = make_lot(w["material"], ids["kg"], w["rloc"], disposition="REJECTED")
    exp = make_lot(w["material"], ids["kg"], w["aloc"], disposition="APPROVED", expiry=date.today() - timedelta(days=1))
    rt = make_lot(w["material"], ids["kg"], w["aloc"], disposition="APPROVED", retest=date.today() - timedelta(days=1))
    ok = make_lot(w["material"], ids["kg"], w["aloc"], disposition="APPROVED")
    s = db.new_session()
    g = lambda i: s.get(MaterialBatch, i)
    for lid, rule in ((rej, "BR-ISS-002"), (exp, "BR-ISS-004"), (rt, "BR-ISS-004")):
        with pytest.raises(BusinessRuleError) as e:
            lots.assert_issuable(s, g(lid))
        assert e.value.rule_id == rule
    lots.assert_issuable(s, g(ok))
    assert lots.derived_status(s, g(exp)) == "EXPIRED" and lots.derived_status(s, g(rt)) == "RETEST_DUE"
    s.close()


def test_quality_hold_blocks_use_and_release_needs_qa_head_signature(w):
    ids = w["ids"]
    lid = make_lot(w["material"], ids["kg"], w["aloc"])
    s = db.new_session()
    lots.assert_issuable(s, s.get(MaterialBatch, lid))
    s.close()
    assert w["wh1"].post("/api/v1/holds", headers=w["hwh1"], json={"entity_type": "MATERIAL_BATCH", "record_id": lid, "reason": "suspect"}).status_code == 403
    h = w["qo"].post("/api/v1/holds", headers=w["hqo"], json={"entity_type": "MATERIAL_BATCH", "record_id": lid, "reason": "foreign particles found"})
    assert h.status_code == 201 and h.json()["hold_no"].startswith("HLD-")
    d = w["wh1"].get(f"/api/v1/lots/{lid}").json()
    assert d["status"] == "HOLD" and d["issue_violations"][0]["rule_id"] == "BR-HOLD-001"
    s = db.new_session()
    with pytest.raises(BusinessRuleError) as e:
        lots.assert_issuable(s, s.get(MaterialBatch, lid))
    assert e.value.rule_id == "BR-HOLD-001"
    s.close()
    assert w["qo"].post(f"/api/v1/holds/{h.json()['id']}/release", headers=w["hqo"], json={"password": PW, "reason": "ok"}).status_code == 403
    assert w["qa"].post(f"/api/v1/holds/{h.json()['id']}/release", headers=w["hqa"], json={"password": "bad-Password-1!", "reason": "ok"}).status_code == 401
    r = w["qa"].post(f"/api/v1/holds/{h.json()['id']}/release", headers=w["hqa"], json={"password": PW, "reason": "investigation closed, no impact"})
    assert r.status_code == 200 and r.json()["status"] == "RELEASED" and r.json()["release_signature_id"]
    assert w["wh1"].get(f"/api/v1/lots/{lid}").json()["status"] == "APPROVED"


# ------------------------------------------------------------------ transfers, ledger, concurrency
def test_transfer_rules_for_quarantine_approved_and_rejected_stock(w):
    lot = received_lot(w)["lots"][0]
    tr = lambda to, qty=10, **kw: w["wh1"].post(f"/api/v1/lots/{lot['id']}/transfer", headers=w["hwh1"],
                                                json={"from_location_id": w["qloc"], "to_location_id": to, "quantity": qty, "reason": "move", **kw})
    r = tr(w["aloc"])
    assert r.status_code == 409 and r.json()["rule_id"] == "BR-QRN-001"          # unreleased -> quarantine areas only
    ok = tr(w["qloc2"], 30)
    assert ok.status_code == 200 and ok.json()["txn_type"] == "TRANSFER"
    d = w["wh1"].get(f"/api/v1/lots/{lot['id']}").json()
    assert sorted((b["location"], b["on_hand"]) for b in d["balances"]) == [("QZ1", 70.0), ("QZ2", 30.0)]
    assert tr(w["qloc2"], 500).json()["rule_id"] == "BR-INV-002"                  # no negative stock


def test_ledger_matches_balances_and_detects_drift(w, engine):
    received_lot(w)
    s = db.new_session()
    assert inventory.verify_ledger(s) == []
    s.close()
    raw("UPDATE inventory_balance SET qty_on_hand = qty_on_hand + 5")             # someone edits the projection directly
    out = w["qa"].get("/api/v1/inventory/verify").json()
    assert out["ok"] is False and out["discrepancies"][0]["balance"] - out["discrepancies"][0]["ledger"] == 5.0


def test_ledger_is_append_only_and_lot_is_immutable(w):
    lot = received_lot(w)["lots"][0]
    s = db.new_session()
    t = s.execute(select(InventoryTransaction)).scalars().first()
    t.quantity = 1
    with pytest.raises(ImmutableRecordError):
        s.flush()
    s.rollback()
    l = s.get(MaterialBatch, lot["id"])
    l.quantity = 5
    with pytest.raises(ImmutableRecordError):
        s.flush()
    s.rollback()
    with pytest.raises(Exception):
        s.execute(text("DELETE FROM inventory_transaction"))
    s.close()


def test_concurrent_issues_cannot_overdraw_stock(w, engine):
    lid = make_lot(w["material"], w["ids"]["kg"], w["aloc"], qty=10)
    results = []

    def worker():
        s = db.new_session()
        try:
            lot = s.get(MaterialBatch, lid)
            inventory.post(s, txn_type="ISSUE", batch=lot, quantity=Decimal(6), from_location_id=w["aloc"], ref_doc_type="TEST")
            s.commit()
            results.append("ok")
        except BusinessRuleError:
            s.rollback()
            results.append("blocked")
        except Exception as e:  # noqa: BLE001
            s.rollback()
            results.append(type(e).__name__)
        finally:
            s.close()

    from app.audit.context import AuditContext, audit_context
    ts = [threading.Thread(target=lambda: audit_context(AuditContext(user_name="t", reason="r")) and worker()) for _ in range(2)]
    [t.start() for t in ts]
    [t.join() for t in ts]
    assert sorted(results).count("ok") == 1 and "blocked" in results or sorted(results) == ["ok", "OperationalError"]
    s = db.new_session()
    assert inventory.on_hand(s, lid) == Decimal(4) and inventory.verify_ledger(s) == []
    s.close()


# ------------------------------------------------------------------ FEFO / FIFO
def test_fefo_pick_order_and_exclusions(w):
    k = w["ids"]["kg"]
    d = date.today()
    far = make_lot(w["material"], k, w["aloc"], qty=50, expiry=d + timedelta(days=400))
    near = make_lot(w["material"], k, w["aloc"], qty=30, expiry=d + timedelta(days=60))
    mid = make_lot(w["material"], k, w["aloc"], qty=40, expiry=d + timedelta(days=200))
    make_lot(w["material"], k, w["aloc"], qty=99, expiry=d - timedelta(days=1))                 # expired: excluded
    make_lot(w["material"], k, w["qloc"], qty=99, disposition="QUARANTINE")                      # unreleased: excluded
    held = make_lot(w["material"], k, w["aloc"], qty=99, expiry=d + timedelta(days=10))
    w["qo"].post("/api/v1/holds", headers=w["hqo"], json={"entity_type": "MATERIAL_BATCH", "record_id": held, "reason": "under investigation"})
    pick = w["wh1"].get("/api/v1/inventory/fefo-pick", params={"material_id": w["material"], "quantity": 80}).json()
    assert [p["lot_id"] for p in pick] == [near, mid, far] and [p["take"] for p in pick] == [30.0, 40.0, 10.0]
    big = w["wh1"].get("/api/v1/inventory/fefo-pick", params={"material_id": w["material"], "quantity": 500})
    assert big.status_code == 409 and big.json()["rule_id"] == "BR-INV-002"


# ------------------------------------------------------------------ labels
def test_labels_are_controlled_audited_and_limited(w):
    lot = received_lot(w)["lots"][0]
    r = w["wh1"].post(f"/api/v1/lots/{lot['id']}/labels", headers=w["hwh1"], json={"label_type": "QUARANTINE", "copies": 3})
    assert r.status_code == 200 and r.content.startswith(b"%PDF") and r.headers["x-label-no"].startswith("LBL-")
    assert r.content.count(b"/Type /Page\n") >= 3 or r.content.count(b"/Type /Page") >= 3
    again = w["wh1"].post(f"/api/v1/lots/{lot['id']}/labels", headers=w["hwh1"], json={"label_type": "QUARANTINE", "copies": 1})
    assert again.status_code == 422 and again.json()["code"] == "REASON_REQUIRED"            # reprint needs a reason
    assert w["wh1"].post(f"/api/v1/lots/{lot['id']}/labels", headers=w["hwh1"], json={"label_type": "QUARANTINE", "copies": 1, "reprint_reason": "label damaged"}).status_code == 200
    assert w["wh1"].post(f"/api/v1/lots/{lot['id']}/labels", headers=w["hwh1"], json={"label_type": "QUARANTINE", "copies": 11, "reprint_reason": "x"}).status_code == 422
    blocked = w["wh1"].post(f"/api/v1/lots/{lot['id']}/labels", headers=w["hwh1"], json={"label_type": "APPROVED", "copies": 1})
    assert blocked.status_code == 409 and blocked.json()["rule_id"] == "BR-LBL-002"          # not released
    hist = w["wh1"].get(f"/api/v1/lots/{lot['id']}/labels").json()
    assert len(hist) == 2 and hist[1]["reprint_reason"] == "label damaged"
    s = db.new_session()
    assert s.execute(select(AuditTrail).where(AuditTrail.action == "LABEL_PRINT")).scalars().all().__len__() == 2
    s.close()
    ap = make_lot(w["material"], w["ids"]["kg"], w["aloc"])
    assert w["wh1"].post(f"/api/v1/lots/{ap}/labels", headers=w["hwh1"], json={"label_type": "APPROVED", "copies": 2}).status_code == 200
    assert w["wh1"].post(f"/api/v1/locations/{w['aloc']}/label", headers=w["hwh1"]).content.startswith(b"%PDF")


# ------------------------------------------------------------------ temperature, destruction, jobs
def test_temperature_excursion_places_holds_on_stored_lots(w):
    lid = make_lot(w["material"], w["ids"]["kg"], w["aloc"])
    ok = w["wh1"].post(f"/api/v1/locations/{w['aloc']}/temperature", headers=w["hwh1"], json={"reading": 20})
    assert ok.json()["excursion"] is False
    bad = w["wh1"].post(f"/api/v1/locations/{w['aloc']}/temperature", headers=w["hwh1"], json={"reading": 31.5, "remarks": "chiller failure"})
    assert bad.status_code == 201 and bad.json()["excursion"] is True and len(bad.json()["holds"]) == 1
    assert w["wh1"].get(f"/api/v1/lots/{lid}").json()["status"] == "HOLD"
    s = db.new_session()
    assert s.execute(select(Notification).where(Notification.category == "TEMPERATURE")).first()
    s.close()
    assert len(w["wh1"].get(f"/api/v1/locations/{w['aloc']}/temperature").json()) == 2


def test_destruction_requires_rejected_stock_and_qa_signature(w):
    ids = w["ids"]
    rej = make_lot(w["material"], ids["kg"], w["rloc"], qty=20, disposition="REJECTED")
    ok_lot = make_lot(w["material"], ids["kg"], w["aloc"], qty=20, disposition="APPROVED")
    body = lambda lid, loc: {"material_batch_id": lid, "location_id": loc, "quantity": 20, "method": "Incineration", "reason": "rejected lot disposal"}
    assert w["wh1"].post("/api/v1/destructions", headers=w["hwh1"], json=body(ok_lot, w["aloc"])).status_code == 409
    d = w["wh1"].post("/api/v1/destructions", headers=w["hwh1"], json=body(rej, w["rloc"])).json()
    assert w["wh1"].post(f"/api/v1/destructions/{d['id']}/execute", headers=w["hwh1"]).status_code == 409     # not approved yet
    assert w["wh2"].post(f"/api/v1/destructions/{d['id']}/decision", headers=w["hwh2"], json={"password": PW, "reason": "x"}).status_code == 403
    assert w["qa"].post(f"/api/v1/destructions/{d['id']}/decision", headers=w["hqa"], json={"password": PW, "reason": "approved for incineration"}).json()["status"] == "APPROVED"
    ex = w["wh1"].post(f"/api/v1/destructions/{d['id']}/execute", headers=w["hwh1"])
    assert ex.status_code == 200 and ex.json()["status"] == "EXECUTED"
    lot = w["wh1"].get(f"/api/v1/lots/{rej}").json()
    assert lot["disposition"] == "DESTROYED" and lot["on_hand"] == 0 and lot["reconciliation"]["destroyed"] == 20.0


def test_expiry_job_and_alerts(w):
    k = w["ids"]["kg"]
    gone = make_lot(w["material"], k, w["aloc"], expiry=date.today() - timedelta(days=3))
    soon = make_lot(w["material"], k, w["aloc"], expiry=date.today() + timedelta(days=20))
    ad, ha = as_user(w["wh1"].app, "admin_b", [])
    r = ad.post("/api/v1/jobs/inventory-expiry", headers=ha)
    assert r.status_code == 200 and r.json()["expired"] == [gone] and r.json()["alerts_created"] >= 1
    assert w["wh1"].get(f"/api/v1/lots/{gone}").json()["disposition"] == "EXPIRED"
    assert ad.post("/api/v1/jobs/inventory-expiry", headers=ha).json()["expired"] == []
    s = db.new_session()
    titles = [n.title for n in s.execute(select(Notification).where(Notification.category == "EXPIRY")).scalars()]
    assert any("expired" in t.lower() for t in titles) and any("within 30 days" in t for t in titles)
    s.close()
    # expired stock may only go to the rejected area
    mv = w["wh1"].post(f"/api/v1/lots/{gone}/transfer", headers=w["hwh1"], json={"from_location_id": w["aloc"], "to_location_id": w["qloc"], "quantity": 1, "reason": "x"})
    assert mv.status_code == 409 and mv.json()["rule_id"] == "BR-QRN-003"


def test_stock_report_and_export(w):
    import io
    from openpyxl import load_workbook
    received_lot(w)
    make_lot(w["material"], w["ids"]["kg"], w["aloc"], qty=7)
    st = w["wh1"].get("/api/v1/inventory/stock").json()
    assert st["total"] == 2 and {r["status"] for r in st["items"]} == {"QUARANTINE", "APPROVED"}
    assert w["wh1"].get("/api/v1/inventory/stock", params={"status": "APPROVED"}).json()["total"] == 1
    x = w["wh1"].get("/api/v1/inventory/stock/export")
    text_ = " ".join(str(c) for row in load_workbook(io.BytesIO(x.content)).active.iter_rows(values_only=True) for c in row if c)
    assert "Stock by lot and location" in text_ and "QZ1" in text_


def test_roles_cannot_do_warehouse_actions_they_do_not_hold(w):
    assert w["pu"].post(GRN, headers=w["hpu"], json=grn_body(w)).status_code == 403
    assert w["qc"].post("/api/v1/inventory/verify", headers=w["hqc"]).status_code in (403, 404, 405)
    assert w["qc"].post(f"/api/v1/destructions", headers=w["hqc"], json={"material_batch_id": 1, "location_id": 1, "quantity": 1, "method": "xxx", "reason": "xxx"}).status_code == 403
    assert w["wh1"].get("/api/v1/inventory/verify").status_code == 403           # ledger verification is QA's
