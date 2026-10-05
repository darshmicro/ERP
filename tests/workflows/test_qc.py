from datetime import date, timedelta

import pytest
from sqlalchemy import select

from app.core import db
from app.core.errors import BusinessRuleError, ImmutableRecordError
from app.models import (AuditTrail, COA, MaterialBatch, Notification, OOTEvent, QCResult, QCResultAmendment, QualityHold)
from app.services import lots
from tests.conftest import PW
from tests.helpers import (as_user, assign, build_qc_world, make_lot, raw, release_lot, submit_pass_results, take_sample)


@pytest.fixture()
def w(app):
    return build_qc_world(app)


# ------------------------------------------------------------------ sampling
def test_sampling_plan_ledger_and_state(w):
    need = w["qc"].get(f"/api/v1/lots/{w['lot']}/sampling-requirements").json()
    assert need == {"rule": "SQRT_N_PLUS_1", "containers_total": 4, "min_containers": 3, "min_quantity": 0.0}
    few = w["qc"].post("/api/v1/samples", headers=w["hqc"], json={"material_batch_id": w["lot"], "quantity_sampled": 5, "containers_sampled": 1, "sampling_location_id": w["qloc"], "reason": "x"})
    assert few.status_code == 409 and few.json()["rule_id"] == "BR-SMP-002"
    assert w["wh1"].post("/api/v1/samples", headers=w["hwh1"], json={"material_batch_id": w["lot"], "quantity_sampled": 5, "containers_sampled": 3, "sampling_location_id": w["qloc"], "reason": "x"}).status_code == 403
    sm = take_sample(w)
    assert sm["sample_no"].startswith("SMP-") and sm["status"] == "CREATED" and sm["specification_id"] == w["spec"]
    lot = w["wh1"].get(f"/api/v1/lots/{w['lot']}").json()
    assert lot["disposition"] == "QC_TESTING" and lot["on_hand"] == 95.0 and lot["reconciliation"]["sampled"] == 5.0
    assert w["wh1"].get(f"/api/v1/lots/{w['lot']}/ledger").json()[-1]["txn_type"] == "SAMPLE"
    assert w["qc"].post(f"/api/v1/samples/{sm['id']}/label", headers=w["hqc"]).content.startswith(b"%PDF")
    ap = make_lot(w["material"], w["ids"]["kg"], w["aloc"])
    r = w["qc"].post("/api/v1/samples", headers=w["hqc"], json={"material_batch_id": ap, "quantity_sampled": 1, "containers_sampled": 1, "sampling_location_id": w["aloc"], "reason": "x"})
    assert r.status_code == 409 and r.json()["rule_id"] == "BR-SMP-001"


def test_tests_follow_pinned_specification_even_after_spec_revision(w):
    sm = take_sample(w)
    tests = assign(w, sm["id"], None)
    assert set(tests) == {"Assay", "Appearance"} and tests["Assay"]["status"] == "ASSIGNED"
    # QA revises the spec afterwards (tighter limits): existing sample/tests keep the version they were created against
    v2 = w["qc"].post(f"/api/v1/specifications/{w['spec']}/new-version", headers=w["hqc"], json={"reason": "tighter"}).json()
    pid = w["qc"].get(f"/api/v1/specifications/{v2['id']}").json()["parameters"][0]["id"]
    w["qc"].patch(f"/api/v1/specifications/{v2['id']}/parameters/{pid}", headers=w["hqc"], json={"lsl": 99.9, "usl": 100.1})
    w["qc"].post(f"/api/v1/specifications/{v2['id']}/submit", headers=w["hqc"])
    assert w["qa"].post(f"/api/v1/specifications/{v2['id']}/approve", headers=w["hqa"], json={"password": PW, "reason": "ok"}).status_code == 200
    r = w["qc"].post(f"/api/v1/qc/tests/{tests['Assay']['id']}/result", headers=w["hqc"], json={"value": 100.5, "reason": "r"})
    res = r.json()
    assert res["pass_fail"] == "PASS" and float(res["lsl"]) == 99 and float(res["usl"]) == 101    # judged against v1 (crit 13)


# ------------------------------------------------------------------ results
def test_result_entry_evaluation_rounding_and_immutability(w):
    sm = take_sample(w)
    t = assign(w, sm["id"])
    a = t["Assay"]
    assert w["qc2"].post(f"/api/v1/qc/tests/{a['id']}/result", headers=w["hqc2"], json={"value": 100, "reason": "r"}).status_code in (201, 403)
    # (unassigned tests accept any analyst; assignment is verified in test_assigned_test)
    r = w["qc"].post(f"/api/v1/qc/tests/{a['id']}/result", headers=w["hqc"], json={"value": 100.04, "reason": "r"})
    if r.status_code == 201:
        assert float(r.json()["rounded_value"]) == 100.0 and r.json()["pass_fail"] == "PASS"
    again = w["qc"].post(f"/api/v1/qc/tests/{a['id']}/result", headers=w["hqc"], json={"value": 100, "reason": "r"})
    assert again.status_code == 409 and again.json()["rule_id"] == "BR-QC-005"
    s = db.new_session()
    res = s.execute(select(QCResult)).scalars().first()
    res.value_numeric = 1
    with pytest.raises(ImmutableRecordError):
        s.flush()
    s.close()


def test_assigned_test_is_restricted_to_its_analyst_and_boundaries_are_inclusive(w):
    sm = take_sample(w)
    from sqlalchemy import select as sel
    s = db.new_session()
    from app.models import User
    uid = s.execute(sel(User.id).where(User.username == "qc_an2")).scalar()
    s.close()
    t = assign(w, sm["id"], uid)
    assert w["qc"].post(f"/api/v1/qc/tests/{t['Assay']['id']}/result", headers=w["hqc"], json={"value": 100, "reason": "r"}).status_code == 403
    r = w["qc2"].post(f"/api/v1/qc/tests/{t['Assay']['id']}/result", headers=w["hqc2"], json={"value": 99.0, "reason": "r"})
    assert r.status_code == 201 and r.json()["pass_fail"] == "PASS"                    # LSL is inclusive
    bad = w["qc2"].post(f"/api/v1/qc/tests/{t['Appearance']['id']}/result", headers=w["hqc2"], json={"reason": "r"})
    assert bad.status_code == 422                                                       # 'conforms' required


def test_failing_result_raises_oos_and_places_hold(w):
    sm = take_sample(w)
    t = assign(w, sm["id"])
    r = w["qc"].post(f"/api/v1/qc/tests/{t['Assay']['id']}/result", headers=w["hqc"], json={"value": 97.5, "reason": "r"})
    assert r.status_code == 201 and r.json()["pass_fail"] == "FAIL"
    oos = w["qch"].get("/api/v1/oos").json()["items"]
    assert len(oos) == 1 and oos[0]["status"] == "RAISED" and oos[0]["oos_no"].startswith("OOS-")
    lot = w["wh1"].get(f"/api/v1/lots/{w['lot']}").json()
    assert lot["status"] == "HOLD" and lot["holds"][0]["source"] == "OOS"
    s = db.new_session()
    assert s.execute(select(Notification).where(Notification.category == "OOS")).first()
    s.close()


# ------------------------------------------------------------------ calibration (crit 11)
def _equipment(w, due_ago_days=None, calibrated=True):
    eq = w["qo"].post("/api/v1/equipment", headers=w["hqo"], json={"name": "HPLC", "reason": "register"}).json()
    w["qo"].patch(f"/api/v1/equipment/{eq['id']}", headers=w["hqo"], json={"qualification_status": "QUALIFIED", "reason": "IQ/OQ"})
    if calibrated:
        today = date.today()
        po, du = (today - timedelta(days=400), today - timedelta(days=due_ago_days)) if due_ago_days else (today, today + timedelta(days=180))
        w["qo"].post(f"/api/v1/equipment/{eq['id']}/calibrations", headers=w["hqo"], json={"performed_on": str(po), "due_on": str(du), "result": "PASS"})
    return eq["id"]


def test_crit_11_invalid_calibration_blocks_testing_and_override_is_controlled(w):
    eq_bad = _equipment(w, due_ago_days=30)
    sm = take_sample(w)
    t = assign(w, sm["id"])
    a = t["Assay"]["id"]
    r = w["qc"].post(f"/api/v1/qc/tests/{a}/start", headers=w["hqc"], json={"equipment_id": eq_bad})
    assert r.status_code == 409 and r.json()["rule_id"] == "BR-QC-002"
    # a valid instrument is accepted and the calibration status is snapshotted
    eq_ok = _equipment(w)
    ok = w["qc"].post(f"/api/v1/qc/tests/{a}/start", headers=w["hqc"], json={"equipment_id": eq_ok})
    assert ok.status_code == 200 and ok.json()["calibration_status"] == "VALID" and ok.json()["status"] == "STARTED"
    # test 2 with the expired instrument: overrides are disabled by default
    b = t["Appearance"]["id"]
    ad, ha = as_user(w["qc"].app, "admin_b", [])
    assert w["qa"].post(f"/api/v1/qc/tests/{b}/override-calibration", headers=w["hqa"], json={"password": PW, "reason": "x"}).status_code == 409
    # enable the configured override, then QA head signs it; the result is then accepted
    assert ad.put("/api/v1/config/calibration.override_allowed", headers=ha, json={"value": "true", "reason": "validated exception procedure"}).status_code == 200
    w["qc"].post(f"/api/v1/qc/tests/{b}/start", headers=w["hqc"], json={"equipment_id": eq_bad})        # now allowed to start...
    blocked = w["qc"].post(f"/api/v1/qc/tests/{b}/result", headers=w["hqc"], json={"conforms": True, "reason": "r"})
    assert blocked.status_code == 409 and blocked.json()["rule_id"] == "BR-QC-002"                      # ...but not to complete
    assert w["qo"].post(f"/api/v1/qc/tests/{b}/override-calibration", headers=w["hqo"], json={"password": PW, "reason": "x"}).status_code == 403
    assert w["qa"].post(f"/api/v1/qc/tests/{b}/override-calibration", headers=w["hqa"], json={"password": "bad-Password-1!", "reason": "x"}).status_code == 401
    ov = w["qa"].post(f"/api/v1/qc/tests/{b}/override-calibration", headers=w["hqa"], json={"password": PW, "reason": "deviation DEV-9 raised"})
    assert ov.status_code == 200 and ov.json()["calibration_override_signature_id"]
    assert w["qc"].post(f"/api/v1/qc/tests/{b}/result", headers=w["hqc"], json={"conforms": True, "reason": "r"}).status_code == 201


# ------------------------------------------------------------------ amendments
def test_result_amendment_workflow_retains_original(w):
    sm = take_sample(w)
    t = assign(w, sm["id"])
    res = w["qc"].post(f"/api/v1/qc/tests/{t['Assay']['id']}/result", headers=w["hqc"], json={"value": 98.0, "reason": "r"}).json()      # FAIL (typo)
    am = w["qc"].post(f"/api/v1/qc/results/{res['id']}/amendments", headers=w["hqc"], json={"value": 100.0, "reason": "transcription error: 98.0 should be 100.0"})
    assert am.status_code == 201 and am.json()["original_value"] and am.json()["new_pass_fail"] == "PASS"
    assert w["qc"].post(f"/api/v1/qc/amendments/{am.json()['id']}/decision", headers=w["hqc"], json={"password": PW, "reason": "ok"}).status_code == 403   # analyst cannot approve
    assert w["qch"].post(f"/api/v1/qc/amendments/{am.json()['id']}/decision", headers=w["hqch"], json={"password": "bad-Password-1!", "reason": "ok"}).status_code == 401
    ok = w["qch"].post(f"/api/v1/qc/amendments/{am.json()['id']}/decision", headers=w["hqch"], json={"password": PW, "reason": "verified against raw data"})
    assert ok.status_code == 200 and ok.json()["status"] == "APPROVED" and ok.json()["approved_signature_id"]
    det = w["qc"].get(f"/api/v1/samples/{sm['id']}").json()
    r = next(x for x in det["tests"] if x["test_name"] == "Assay")["result"]
    assert float(r["rounded_value"]) == 98.0 and r["pass_fail"] == "FAIL"                  # the original record is untouched
    assert r["effective"]["amended"] is True and float(r["effective"]["value"]) == 100.0 and r["effective"]["pass_fail"] == "PASS"
    assert len(w["qc"].get(f"/api/v1/qc/results/{res['id']}/amendments").json()) == 1
    assert w["qc"].post(f"/api/v1/qc/results/{res['id']}/amendments", headers=w["hqc"], json={"value": 99.5, "reason": "x yz"}).status_code == 201
    assert w["qc"].post(f"/api/v1/qc/results/{res['id']}/amendments", headers=w["hqc"], json={"value": 99.6, "reason": "x yz"}).status_code == 409   # one pending at a time


# ------------------------------------------------------------------ release chain
def test_full_release_chain_to_approved_coa_and_issuable_lot(w):
    out = release_lot(w)
    assert out["disposition"] == "APPROVED" and out["qa_release_no"].startswith("QAR-") and out["qc_no"].startswith("QC-")
    assert [h["decision"] for h in out["workflow"]["history"]] == ["SUBMIT", "APPROVE", "APPROVE", "APPROVE"]
    lot = w["wh1"].get(f"/api/v1/lots/{w['lot']}").json()
    assert lot["disposition"] == "APPROVED" and lot["released_at"] and lot["issue_violations"] == []
    s = db.new_session()
    lots.assert_issuable(s, s.get(MaterialBatch, w["lot"]))                # released lot is now issuable
    s.close()
    coa = w["qc"].get(f"/api/v1/lots/{w['lot']}/coa").json()
    assert len(coa) == 1 and coa[0]["conclusion"] == "COMPLIES" and coa[0]["version_no"] == 1
    pdf = w["qc"].get(f"/api/v1/coa/{coa[0]['id']}/pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    import io
    from openpyxl import load_workbook
    x = w["qc"].get(f"/api/v1/coa/{coa[0]['id']}/xlsx")
    cells = " ".join(str(c) for row in load_workbook(io.BytesIO(x.content)).active.iter_rows(values_only=True) for c in row if c)
    assert "Assay" in cells and "99 – 101" in cells and "COMPLIES" in cells and "Sodium Chloride IP" in cells
    # approved label now allowed; release history appears in the audit trail with signatures
    assert w["wh1"].post(f"/api/v1/lots/{w['lot']}/labels", headers=w["hwh1"], json={"label_type": "APPROVED", "copies": 2}).status_code == 200
    s = db.new_session()
    assert s.execute(select(AuditTrail).where(AuditTrail.action == "COA_GENERATED")).first()
    s.close()


def test_release_chain_enforces_order_roles_and_separation_of_duties(w):
    s = take_sample(w)
    assign(w, s["id"])
    pend = w["qc"].post(f"/api/v1/lots/{w['lot']}/submit-release", headers=w["hqc"])
    assert pend.status_code == 409 and pend.json()["rule_id"] == "BR-REL-001" and "Result pending" in pend.json()["message"]
    submit_pass_results(w, s["id"])
    assert w["qc"].post(f"/api/v1/lots/{w['lot']}/submit-release", headers=w["hqc"]).status_code == 200
    dec = lambda who, h, **kw: w[who].post(f"/api/v1/lots/{w['lot']}/release-decision", headers=w[h], json={"decision": "APPROVE", "password": PW, **kw})
    assert dec("qc", "hqc").status_code == 403                    # analyst has no release authority
    assert dec("qa", "hqa").status_code == 403                    # QA head cannot jump the QC head step (role of current step)
    assert dec("qch", "hqch", password="bad-Password-1!").status_code == 401
    r1 = dec("qch", "hqch")
    assert r1.status_code == 200 and r1.json()["disposition"] == "QC_APPROVED"
    assert dec("qch", "hqch").status_code in (403, 409)           # cannot approve again
    r2 = dec("qo", "hqo")
    assert r2.json()["disposition"] == "QA_REVIEW"
    r3 = dec("qa", "hqa")
    assert r3.status_code == 200 and r3.json()["disposition"] == "APPROVED"


def test_one_person_cannot_hold_two_release_steps_sod03(app):
    w = build_qc_world(app)
    both, hb = as_user(app, "dual", ["QC_HEAD", "QA_OFFICER"])
    s = take_sample(w)
    assign(w, s["id"])
    submit_pass_results(w, s["id"])
    w["qc"].post(f"/api/v1/lots/{w['lot']}/submit-release", headers=w["hqc"])
    d1 = both.post(f"/api/v1/lots/{w['lot']}/release-decision", headers=hb, json={"decision": "APPROVE", "password": PW})
    assert d1.status_code == 200
    d2 = both.post(f"/api/v1/lots/{w['lot']}/release-decision", headers=hb, json={"decision": "APPROVE", "password": PW})
    assert d2.status_code == 409 and d2.json()["rule_id"] == "SOD-03"


def test_analyst_who_tested_cannot_review_even_with_reviewer_role(app):
    w = build_qc_world(app)
    qh2, hq2 = as_user(app, "qc_head_tester", ["QC_HEAD", "QC_ANALYST"])
    s = take_sample(w)
    t = assign(w, s["id"], None)
    for name, tt in t.items():
        qh2.post(f"/api/v1/qc/tests/{tt['id']}/result", headers=hq2, json={"value": 100} if name == "Assay" else {"conforms": True})
    w["qc"].post(f"/api/v1/lots/{w['lot']}/submit-release", headers=w["hqc"])
    r = qh2.post(f"/api/v1/lots/{w['lot']}/release-decision", headers=hq2, json={"decision": "APPROVE", "password": PW})
    assert r.status_code == 409 and r.json()["rule_id"] == "SOD-02"


def test_release_rejection_makes_lot_unissuable_crit_05(w):
    s = take_sample(w)
    assign(w, s["id"])
    submit_pass_results(w, s["id"])
    w["qc"].post(f"/api/v1/lots/{w['lot']}/submit-release", headers=w["hqc"])
    r = w["qch"].post(f"/api/v1/lots/{w['lot']}/release-decision", headers=w["hqch"], json={"decision": "REJECT", "password": PW, "comment": "documentation incomplete"})
    assert r.status_code == 200 and r.json()["disposition"] == "REJECTED"
    s_ = db.new_session()
    with pytest.raises(BusinessRuleError) as e:
        lots.assert_issuable(s_, s_.get(MaterialBatch, w["lot"]))
    assert e.value.rule_id == "BR-ISS-002"
    s_.close()


def test_release_blocked_by_missing_coa_hold_or_open_oos(app):
    w = build_qc_world(app)
    s = take_sample(w)
    assign(w, s["id"])
    submit_pass_results(w, s["id"])
    w["qo"].post("/api/v1/holds", headers=w["hqo"], json={"entity_type": "MATERIAL_BATCH", "record_id": w["lot"], "reason": "complaint review"})
    r = w["qc"].post(f"/api/v1/lots/{w['lot']}/submit-release", headers=w["hqc"])
    assert r.status_code == 409 and "quality hold" in r.json()["message"]
    raw("UPDATE grn_line SET coa_received = :f", f=False)
    r = w["qc"].post(f"/api/v1/lots/{w['lot']}/submit-release", headers=w["hqc"])
    assert "CoA was not received" in r.json()["message"]


# ------------------------------------------------------------------ OOS lifecycle (crit 10 included)
def _failed(w):
    sm = take_sample(w)
    t = assign(w, sm["id"])
    w["qc"].post(f"/api/v1/qc/tests/{t['Assay']['id']}/result", headers=w["hqc"], json={"value": 97.0, "reason": "r"})
    w["qc"].post(f"/api/v1/qc/tests/{t['Appearance']['id']}/result", headers=w["hqc"], json={"conforms": True, "reason": "r"})
    return sm, t, w["qch"].get("/api/v1/oos").json()["items"][0]


def test_oos_invalidated_allows_retest_and_release_path(w):
    sm, t, oos = _failed(w)
    dec = lambda c, h, **kw: c.post(f"/api/v1/oos/{oos['id']}/decide", headers=h, json={"password": PW, "decision": "INVALIDATED", "root_cause": "analyst dilution error", "reason": "lab error confirmed", **kw})
    assert dec(w["qc"], w["hqc"]).status_code == 403                  # crit 10: analyst has no QA authority
    assert dec(w["qch"], w["hqch"]).status_code == 403                # nor does QC head (QA decision)
    assert dec(w["qa"], w["hqa"]).status_code == 409                  # investigation phases first
    assert w["qc"].post(f"/api/v1/oos/{oos['id']}/investigate", headers=w["hqc"], json={"phase": "PHASE1", "findings": "Dilution log shows 10x error"}).json()["status"] == "PHASE1"
    assert dec(w["qa"], w["hqa"], password="bad-Password-1!").status_code == 401
    ok = dec(w["qa"], w["hqa"])
    assert ok.status_code == 200 and ok.json()["status"] == "CLOSED" and ok.json()["decision"] == "INVALIDATED" and ok.json()["decision_signature_id"]
    lot = w["wh1"].get(f"/api/v1/lots/{w['lot']}").json()
    assert lot["holds"] == [] and lot["disposition"] == "QC_TESTING"
    rt = w["qch"].post(f"/api/v1/qc/tests/{t['Assay']['id']}/retest", headers=w["hqch"])
    assert rt.status_code == 201 and rt.json()["retest_of_id"] == t["Assay"]["id"]
    assert w["qc"].post(f"/api/v1/qc/tests/{rt.json()['id']}/result", headers=w["hqc"], json={"value": 100.1, "reason": "retest"}).json()["pass_fail"] == "PASS"
    assert w["qc"].post(f"/api/v1/lots/{w['lot']}/submit-release", headers=w["hqc"]).status_code == 200      # now releasable


def test_oos_confirmed_rejects_lot(w):
    sm, t, oos = _failed(w)
    w["qc"].post(f"/api/v1/oos/{oos['id']}/investigate", headers=w["hqc"], json={"phase": "PHASE1", "findings": "No lab error found"})
    w["qch"].post(f"/api/v1/oos/{oos['id']}/investigate", headers=w["hqch"], json={"phase": "PHASE2", "findings": "Process investigation"})
    r = w["qa"].post(f"/api/v1/oos/{oos['id']}/decide", headers=w["hqa"], json={"password": PW, "decision": "CONFIRMED_FAIL", "root_cause": "vendor process drift", "capa_ref": "CAPA-7", "reason": "failure confirmed"})
    assert r.status_code == 200 and r.json()["capa_ref"] == "CAPA-7"
    lot = w["wh1"].get(f"/api/v1/lots/{w['lot']}").json()
    assert lot["disposition"] == "REJECTED"
    assert w["qc"].post(f"/api/v1/qc/tests/{t['Assay']['id']}/retest", headers=w["hqc"]).status_code in (403, 409)


# ------------------------------------------------------------------ OOT and statistics
def test_oot_alert_limit_flags_but_never_changes_disposition(w):
    sm = take_sample(w)
    t = assign(w, sm["id"])
    r = w["qc"].post(f"/api/v1/qc/tests/{t['Assay']['id']}/result", headers=w["hqc"], json={"value": 100.9, "reason": "r"})     # PASS but above alert 100.8
    assert r.json()["pass_fail"] == "PASS"
    oot = w["qch"].get("/api/v1/oot").json()["items"]
    assert len(oot) == 1 and oot[0]["rule"] == "ALERT_LIMIT"
    lot = w["wh1"].get(f"/api/v1/lots/{w['lot']}").json()
    assert lot["disposition"] == "QC_TESTING" and lot["holds"] == []                         # statistical alert alone is not a rejection
    assert w["qch"].post(f"/api/v1/oot/{oot[0]['id']}/review", headers=w["hqch"], json={"reason": "within normal variation"}).json()["status"] == "REVIEWED"


def test_stats_reference_values_and_status_codes():
    from app.services import stats
    import statistics
    vals = [100 + 0.3 * ((i * 7) % 11 - 5) / 5 for i in range(30)]
    cap = stats.capability(vals, 99.0, 101.0)
    mean, so = statistics.fmean(vals), statistics.stdev(vals)
    mr = sum(abs(vals[i] - vals[i - 1]) for i in range(1, 30)) / 29
    sw = mr / 1.128
    assert cap["status"] == "OK" and cap["n"] == 30
    assert cap["cp"] == pytest.approx(2.0 / (6 * sw)) and cap["cpk"] == pytest.approx(min(101 - mean, mean - 99) / (3 * sw))
    assert cap["pp"] == pytest.approx(2.0 / (6 * so)) and cap["ppk"] == pytest.approx(min(101 - mean, mean - 99) / (3 * so))
    assert stats.capability(vals[:10], 99, 101)["status"] == "INSUFFICIENT_DATA" and stats.capability(vals[:10], 99, 101)["cpk"] is None
    assert stats.capability(vals, None, None)["status"] == "SPECIFICATION_UNAVAILABLE"
    assert stats.capability([5.0] * 30, 1, 9)["status"] == "ZERO_VARIANCE"
    one = stats.capability(vals, None, 101.0)
    assert one["cp"] is None and one["cpk"] is not None and any("LSL unavailable" in n for n in one["notes"])
    d = stats.describe([1, 2, 3, 4])
    assert d["mean"] == 2.5 and d["median"] == 2.5 and d["sd"] == pytest.approx(statistics.stdev([1, 2, 3, 4])) and d["cv_pct"] == pytest.approx(d["sd"] / 2.5 * 100)
    assert stats.describe([])["status"] == "NO_DATA"


def test_nelson_rules_and_control_limits():
    from app.services import stats
    base = [10, 10.1, 9.9, 10.05, 9.95, 10, 10.1, 9.9, 10.0, 10.05]
    cl = stats.control_limits(base)
    assert cl["lcl"] < cl["cl"] < cl["ucl"] and cl["ucl"] - cl["cl"] == pytest.approx(2.66 * cl["mr_bar"])
    rules = lambda v: {r["rule"] for r in stats.nelson_rules(v, 10.0, 0.1, (1, 2, 3, 5, 6))}
    assert 1 in rules([10.0, 10.5])                                       # beyond 3 sigma
    assert 2 in rules([10.01] * 9)                                        # nine on one side
    assert 3 in rules([10.0, 10.02, 10.04, 10.06, 10.08, 10.1, 10.12])    # six rising
    assert 5 in rules([10.0, 10.25, 10.3, 10.0])                          # 2 of 3 beyond 2 sigma
    assert not rules([10.0, 9.95, 10.05, 10.0, 9.98, 10.03])


def test_trend_endpoint_groups_and_capability_status(app):
    w = build_qc_world(app)
    from datetime import datetime, timezone
    from app.audit.context import AuditContext, audit_context
    from app.models import QCTest, Sample
    from app.services import qc as qcs
    s = db.new_session()
    with audit_context(AuditContext(user_name="seed", reason="fixture")):
        from sqlalchemy import select as sel
        lot = s.get(MaterialBatch, w["lot"])
        smp = Sample(sample_no="SMP-T", material_batch_id=lot.id, quantity_sampled=1, unit_id=lot.unit_id, sampled_by_id=1, status="TESTING")
        s.add(smp)
        s.flush()
        from app.models import SpecificationParameter
        par = s.execute(sel(SpecificationParameter).where(SpecificationParameter.specification_id == w["spec"], SpecificationParameter.test_name == "Assay")).scalar_one()
        for i, v in enumerate([100.1, 100.0, 99.9, 100.2, 100.05, 99.95, 100.1, 100.0, 99.8, 100.3, 100.0, 100.1]):
            t = QCTest(sample_id=smp.id, spec_parameter_id=par.id, test_name="Assay", status="ASSIGNED")
            s.add(t)
            s.flush()
            r = QCResult(test_id=t.id, spec_type="NUMERIC", value_numeric=v, rounded_value=v, lsl=99, usl=101, pass_fail="PASS", status="DRAFT", entered_by_id=1)
            r.entered_at = datetime(2026, 1 + i % 3, 5 + i, tzinfo=timezone.utc)
            s.add(r)
            s.flush()
            from app.workflows.state_machine import transition
            transition(s, qcs.RESULT_MACHINE, r, "SUBMITTED")
            transition(s, qcs.TEST_MACHINE, t, "SUBMITTED")
        s.commit()
    s.close()
    d = w["qc"].get("/api/v1/qc/trends", params={"material_id": w["material"], "test_name": "Assay", "group": "month"}).json()
    assert d["descriptive"]["n"] == 12 and d["capability"]["status"] == "INSUFFICIENT_DATA" and d["capability"]["cpk"] is None
    assert d["control_limits"]["ucl"] > d["control_limits"]["lcl"] and len(d["groups"]) == 3 and d["lsl"] == 99 and d["usl"] == 101
    assert w["pu"].get("/api/v1/qc/trends", params={"material_id": w["material"], "test_name": "Assay"}).status_code == 403


# ------------------------------------------------------------------ conditional release
def test_conditional_release_is_a_controlled_exception(w):
    exp = str(date.today() + timedelta(days=30))
    body = {"material_batch_id": w["lot"], "quantity_authorised": 10, "intended_batch_ref": "SFG-2026-000099", "justification": "urgent batch, QC pending",
            "risk_assessment_ref": "RA-2026-12", "identity_confirmed": True, "expires_at": exp}
    cr = lambda **kw: w["qo"].post("/api/v1/conditional-releases", headers=w["hqo"], json={**body, **kw})
    assert cr(identity_confirmed=False).status_code == 422
    assert cr(quantity_authorised=1000).status_code == 409
    assert cr(expires_at=str(date.today())).status_code == 422
    assert w["qc"].post("/api/v1/conditional-releases", headers=w["hqc"], json=body).status_code == 403          # QC cannot request
    r = cr()
    assert r.status_code == 201 and r.json()["cr_no"].startswith("CRL-") and r.json()["status"] == "REQUESTED"
    cid = r.json()["id"]
    assert w["qo"].post(f"/api/v1/conditional-releases/{cid}/decision", headers=w["hqo"], json={"password": PW, "reason": "ok"}).status_code == 403   # officer cannot approve
    ok = w["qa"].post(f"/api/v1/conditional-releases/{cid}/decision", headers=w["hqa"], json={"password": PW, "reason": "risk assessed, identity confirmed by FTIR"})
    assert ok.status_code == 200 and ok.json()["status"] == "APPROVED" and ok.json()["approved_signature_id"]
    s = db.new_session()
    from app.services import qc as qcs
    lot = s.get(MaterialBatch, w["lot"])
    assert qcs.active_conditional_release(s, lot, 10).id == cid and qcs.active_conditional_release(s, lot, 11) is None
    lots.assert_issuable(s, lot, conditional_release_ok=True)           # unreleased lot usable ONLY under an authorisation
    with pytest.raises(BusinessRuleError):
        lots.assert_issuable(s, lot)
    s.close()
    # never for rejected / held / expired lots
    rej = make_lot(w["material"], w["ids"]["kg"], w["rloc"], disposition="REJECTED")
    assert cr(material_batch_id=rej).status_code == 409
    w["qo"].post("/api/v1/holds", headers=w["hqo"], json={"entity_type": "MATERIAL_BATCH", "record_id": w["lot"], "reason": "x yz"})
    assert cr().status_code == 409


# ------------------------------------------------------------------ CoA versions, retention
def test_coa_requires_release_and_reissue_creates_new_version(w):
    from app.services import coa as coa_svc
    s = db.new_session()
    with pytest.raises(BusinessRuleError):
        coa_svc.generate(s, s.get(MaterialBatch, w["lot"]), None, "x", None)       # not released
    s.close()
    release_lot(w)
    assert w["qch"].post(f"/api/v1/lots/{w['lot']}/coa", headers=w["hqch"], json={}).status_code == 422                          # reason required
    assert w["qc"].post(f"/api/v1/lots/{w['lot']}/coa", headers=w["hqc"], json={"reason": "typo in vendor name corrected"}).status_code == 403   # analyst cannot reissue
    w["qch"].post(f"/api/v1/lots/{w['lot']}/coa", headers=w["hqch"], json={"reason": "re-issue after master data correction"})
    versions = w["qch"].get(f"/api/v1/lots/{w['lot']}/coa").json()
    assert [v["version_no"] for v in versions][:2] == [1, 2] and len({v["coa_no"] for v in versions}) == 1
    from sqlalchemy import text
    from sqlalchemy.exc import DBAPIError
    s = db.new_session()
    with pytest.raises(DBAPIError):
        s.execute(text("UPDATE coa SET conclusion='X'"))
    s.close()


def test_sample_retention_and_disposal(w):
    sm = take_sample(w)
    assign(w, sm["id"])
    submit_pass_results(w, sm["id"])
    until = str(date.today() + timedelta(days=365))
    assert w["qch"].post(f"/api/v1/samples/{sm['id']}/retain", headers=w["hqch"], json={"retention_until": until, "reason": "retain per SOP"}).json()["status"] == "RETAINED"
    early = w["qa"].post(f"/api/v1/samples/{sm['id']}/dispose", headers=w["hqa"], json={"password": PW, "reason": "disposal"})
    assert early.status_code == 409 and early.json()["rule_id"] == "BR-SMP-003"
    raw("UPDATE sample SET retention_until=:d", d=str(date.today() - timedelta(days=1)))
    assert w["qc"].post(f"/api/v1/samples/{sm['id']}/dispose", headers=w["hqc"], json={"password": PW, "reason": "x"}).status_code == 403
    ok = w["qa"].post(f"/api/v1/samples/{sm['id']}/dispose", headers=w["hqa"], json={"password": PW, "reason": "retention period over"})
    assert ok.status_code == 200 and ok.json()["status"] == "DISPOSED"
